"""CV Automation — multi-image queue with per-image settings."""
from __future__ import annotations
import os, re, csv, json, hashlib, tempfile, io, zipfile, copy, struct, random, math, base64
from datetime import datetime, timedelta
import shutil, subprocess
from io import BytesIO
from pathlib import Path
import streamlit as st
import streamlit.components.v1 as components
from PIL import Image, ImageDraw
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, RGBColor, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from cv_core.image_processing import auto_detect_boxes, clamp_box, default_boxes, trim_whitespace
from cv_core.drag_crop import drag_crop_box

# ----------------------------------------------------------------------
# Constants
# ----------------------------------------------------------------------
BASE_DIR = Path(__file__).resolve().parent
REAL_AGENCIES = ["Asail", "Al Zaid"]
AGENCY_FOLDER = {"Asail": "Asail", "Al Zaid": "AlZaid"}
BOTH_LABEL = "Both"
TEMPLATE_CANDIDATES = {
    "Non-Experienced": ["Non Experianced.docx", "Non_Experianced.docx"],
    "Experienced": ["Experianced.docx"],
}
EXPERIENCE_LEVELS = ["Non-Experienced", "Experienced"]
VALIDITY_OPTIONS = [5, 10]
RELIGION_OPTIONS = ["MUSLIM", "NON MUSLIM"]
DEFAULT_POSITION = "HOUSE MAID"
COUNTRY_PRESETS = ["Saudi Arabia", "Indonesia", "SUDAN", "Beirut Lebanon", "Jordan", "Kuwait", "Dubai UAE"]
YEARS_RANGE = list(range(1, 15))
OTHER_LABEL = "Other (type manually)"

CROP_KEYS = ["passport", "3x4", "full", "ocr"]
CROP_LABELS = {"passport": "Passport page", "3x4": "3×4 photo", "full": "Full body photo", "ocr": "OCR reading zone"}
CROP_HINTS = {
    "passport": "The full passport bio page — used for the CV's passport photo.",
    "3x4": "Just the small headshot photo.",
    "full": "The full-length/standing photo.",
    "ocr": "Tightest possible crop around the printed text — improves OCR accuracy.",
}
CROP_COLORS = {"passport": "#1976d2", "3x4": "#f57c00", "full": "#00897b", "ocr": "#8e24aa"}

BOX_3X4_IN = (1.10, 1.30); BOX_FULL_IN = (1.85, 3.30); BOX_PASSPORT_BASE_IN = (1.38, 1.59)
DEFAULT_K_3X4 = 1.0; DEFAULT_K_FULL = 1.0; DEFAULT_K_PASSPORT = 3.0
IMAGE_RIGHT_SHIFT_IN = 0.15
K_OPTIONS = [0.8, 1.0, 1.5, 2.0, 2.5, 3.0, 3.5, 4.0]
K_OPTIONS_SMALL = [0.5, 0.6, 0.7, 0.8, 0.9, 1.0, 1.1, 1.2, 1.3, 1.4, 1.5]
K_OPTIONS_FULL = [0.5, 0.6, 0.7, 0.8, 0.9, 1.0, 1.1, 1.2, 1.3, 1.4, 1.5, 1.6]
FIT_MODE = {"IMAGE_3X4": "contain", "IMAGE_FULL": "contain", "IMAGE_PASSPORT": "contain"}
BOLD_RED_KEYS = {"RELIGION", "YEARS_EXP2"}
CROP_PREVIEW_WIDTH = 700
DEFAULT_NAME_TOKENS = ("PQETH", "PPETH")
PDF_ONLY_AGENCIES_WHEN_BOTH = {"Al Zaid"}
MONTH_OPTIONS = ["Month","Jan-1","Feb-2","Mar-3","Apr-4","May-5","Jun-6","Jul-7","Aug-8","Sep-9","Oct-10","Nov-11","Dec-12"]

IMAGE_WIDGET_KEYS = [
    "dob_mode_radio","issue_mode_radio","expiry_mode_radio","full_name_in","passport_no_in",
    "place_of_birth_in","relative_name","marital_select","children_select","cfg_agency","cfg_level",
    "cfg_validity","cfg_religion","exp_position_in","exp_country_select","exp_country_manual",
    "exp_years_exp","img_notes_input",
]
UPLOADER_KEY = "composite_uploader"
CACHE_KEY = "_uploaded_cache"
HASH_KEY = "_uploaded_hashes"
PREVIEW_KEY = "_preview_state"
SOUND_OPTIONS = ["White noise", "Pink noise", "Brown noise", "Ocean waves", "Rain", "Jungle"]

# ----------------------------------------------------------------------
# Ambient noise generator
# ----------------------------------------------------------------------
@st.cache_data(show_spinner=False)
def _ambient_wav_bytes(kind="Ocean waves", seconds=20, sr=22050):
    n = int(seconds * sr)
    samples = [0.0] * n

    if kind == "White noise":
        for i in range(n):
            samples[i] = random.uniform(-1.0, 1.0) * 0.30
    elif kind == "Pink noise":
        b0 = b1 = b2 = b3 = b4 = b5 = b6 = 0.0
        for i in range(n):
            w = random.uniform(-1.0, 1.0)
            b0 = 0.99886 * b0 + w * 0.0555179
            b1 = 0.99332 * b1 + w * 0.0750759
            b2 = 0.96900 * b2 + w * 0.1538520
            b3 = 0.86650 * b3 + w * 0.3104856
            b4 = 0.55000 * b4 + w * 0.5329522
            b5 = -0.7616 * b5 - w * 0.0168980
            pink = (b0 + b1 + b2 + b3 + b4 + b5 + b6 + w * 0.5362) * 0.11
            b6 = w * 0.115926
            samples[i] = max(-1.0, min(1.0, pink * 1.1))
    elif kind == "Brown noise":
        val = 0.0
        for i in range(n):
            val = 0.995 * val + 0.02 * random.uniform(-1.0, 1.0)
            samples[i] = max(-1.0, min(1.0, val * 4.0)) * 0.65
    elif kind == "Ocean waves":
        val = 0.0
        for i in range(n):
            val = 0.995 * val + 0.02 * random.uniform(-1.0, 1.0)
            t = i / sr
            env = 0.45 + 0.55 * (0.5 + 0.5 * math.sin(2 * math.pi * t / 4.2))
            env *= 0.7 + 0.3 * math.sin(2 * math.pi * t / 0.7)
            samples[i] = max(-1.0, min(1.0, val * 5.0)) * env * 0.75
    elif kind == "Rain":
        val = 0.0
        for i in range(n):
            w = random.uniform(-1.0, 1.0)
            val = 0.85 * val + 0.15 * w
            if random.random() < 0.0009:
                val += random.uniform(-0.55, 0.55)
            samples[i] = max(-1.0, min(1.0, val)) * 0.55
    elif kind == "Jungle":
        val = 0.0
        i = 0
        while i < n:
            val = 0.985 * val + 0.03 * random.uniform(-1.0, 1.0)
            samples[i] = max(-1.0, min(1.0, val * 3.0)) * 0.22
            i += 1
            if i < n and random.random() < 0.0022:
                freq = random.uniform(1500, 4200)
                dur = int(sr * random.uniform(0.06, 0.16))
                for k in range(dur):
                    if i + k >= n: break
                    env = (1.0 - k / dur) ** 2
                    samples[i + k] += math.sin(2 * math.pi * freq * k / sr) * env * 0.20
                i += dur

    edge = int(sr * 0.5)
    for k in range(edge):
        f = k / edge
        samples[k] *= f
        samples[n - 1 - k] *= f

    raw = bytearray()
    for s in samples:
        v = int(max(-1.0, min(1.0, s)) * 28000)
        raw += struct.pack("<h", v)

    header = (b"RIFF" + struct.pack("<I", 36 + len(raw)) + b"WAVE"
              + b"fmt " + struct.pack("<IHHIIHH", 16, 1, 1, sr, sr * 2, 2, 16)
              + b"data" + struct.pack("<I", len(raw)))
    return header + bytes(raw)

@st.cache_data(show_spinner=False)
def _ambient_wav_html(kind="Ocean waves", seconds=20, volume=0.6):
    wav = _ambient_wav_bytes(kind, seconds)
    b64 = base64.b64encode(wav).decode("ascii")
    src = f"data:audio/wav;base64,{b64}"
    return f"""
    <audio id="amb-loop" loop controls style="width:100%; outline:none;"
           onloadeddata="this.volume={volume};">
      <source src="{src}" type="audio/wav">
    </audio>
    """

class _CachedUpload:
    def __init__(self, name, data):
        self.name = name
        self._data = data
        self.size = len(data)
    def getvalue(self): return self._data
    def read(self, *a, **k): return self._data

def _file_hash(b): return hashlib.md5(b).hexdigest()[:10]

def _name_from_filename(filename):
    if not filename: return ""
    base = Path(filename).stem
    name = re.sub(r"[_\-\.]+", " ", base)
    name = re.sub(r"[^A-Za-z\s]", "", name)
    name = re.sub(r"\s+", " ", name).strip().upper()
    if len(name) < 4: return ""
    if name.lower() in ("image","img","photo","screenshot","scan"): return ""
    return name

# ----------------------------------------------------------------------
# Session save / load
# ----------------------------------------------------------------------
def _session_to_json() -> str:
    per = {}
    for name, slot in _per_image_dict().items():
        s = dict(slot)
        if s.get("boxes"):
            s["boxes"] = {k: (list(v) if v else None) for k, v in s["boxes"].items()}
        per[name] = s
    data = {
        "per_image": per,
        "k_3x4": st.session_state.get("k_3x4"),
        "k_full": st.session_state.get("k_full"),
        "k_passport": st.session_state.get("k_passport"),
        "theme_mode": st.session_state.get("theme_mode"),
        "uploaded_hashes": st.session_state.get(HASH_KEY, {}),
    }
    return json.dumps(data, indent=2, default=str)

def _load_session_json(json_str: str):
    data = json.loads(json_str)
    new_per = {}
    for name, s in (data.get("per_image") or {}).items():
        slot = dict(s)
        boxes = slot.get("boxes")
        if isinstance(boxes, dict):
            new_boxes = {}
            for k, v in boxes.items():
                if v: new_boxes[k] = tuple(v)
                else: new_boxes[k] = None
            slot["boxes"] = new_boxes
        new_per[name] = slot
    st.session_state.per_image = new_per
    if data.get("k_3x4") is not None: st.session_state.k_3x4 = float(data["k_3x4"])
    if data.get("k_full") is not None: st.session_state.k_full = float(data["k_full"])
    if data.get("k_passport") is not None: st.session_state.k_passport = float(data["k_passport"])
    if data.get("theme_mode"): st.session_state.theme_mode = data["theme_mode"]
    if data.get("uploaded_hashes"): st.session_state[HASH_KEY] = data["uploaded_hashes"]
    # Force a clean switch on next render
    st.session_state.pop("active_file", None)
    st.session_state.pop(PREVIEW_KEY, None)
    st.session_state["_force_date_reseed"] = True

def _preview_pdf_bytes(cv, agency, level, active_name, files_list, sboxes, rotation=0):
    with tempfile.TemporaryDirectory() as td:
        td_path = Path(td)
        a0 = (REAL_AGENCIES if agency == BOTH_LABEL else [agency])[0]
        t0 = find_template(a0, level)
        if t0 is None: return None
        f_obj = next((f for f in files_list if f.name == active_name), None)
        if not f_obj: return None
        try:
            simg = Image.open(BytesIO(f_obj.getvalue())).convert("RGB")
            if rotation: simg = simg.rotate(-rotation, expand=True)
        except Exception: return None
        imgs = {"IMAGE_PASSPORT": td_path/"passport.jpg","IMAGE_3X4": td_path/"3x4.jpg","IMAGE_FULL": td_path/"full.jpg"}
        simg.crop(sboxes["passport"]).save(imgs["IMAGE_PASSPORT"], quality=95)
        simg.crop(sboxes["3x4"]).save(imgs["IMAGE_3X4"], quality=95)
        simg.crop(sboxes["full"]).save(imgs["IMAGE_FULL"], quality=95)
        docx_p = td_path/"preview.docx"; pdf_p = td_path/"preview.pdf"
        fill_cv(t0, cv, imgs, docx_p, pdf_p)
        if pdf_p.exists(): return pdf_p.read_bytes()
    return None

def _per_image_dict():
    if "per_image" not in st.session_state: st.session_state.per_image = {}
    return st.session_state.per_image

def _blank_image_slot():
    return {"boxes": None,"extracted": None,"cv_data": None,"approved": False,
            "agency": REAL_AGENCIES[0],"level": "Non-Experienced","validity_years": 5,
            "religion": "MUSLIM","position": DEFAULT_POSITION,"country": COUNTRY_PRESETS[0],
            "years_exp": 2,"notes": "", "rotation": 0}

def _clear_image_widgets():
    to_delete = []
    for k in list(st.session_state.keys()):
        if k in IMAGE_WIDGET_KEYS: to_delete.append(k)
        elif k.endswith(("_day","_month","_year")) and not k.startswith("_"): to_delete.append(k)
    for k in to_delete: del st.session_state[k]
    for k in ("_prev_dob_mode","_prev_issue_mode","_prev_expiry_mode","_prev_cfg_validity",
              "_prev_cfg_level","_last_dob_for_suggestion","auto_read_pending","auto_read_source"):
        st.session_state.pop(k, None)

def _switch_to(new_name):
    _clear_image_widgets()
    st.session_state.active_file = new_name
    slot = _per_image_dict().get(new_name, {})
    st.session_state.extracted = slot.get("extracted")
    st.session_state.cv_data = slot.get("cv_data")
    st.session_state.approved = slot.get("approved", False)
    st.session_state["_force_date_reseed"] = True

def _name_issue_reason(name):
    if not name or not name.strip(): return ""
    n = name.strip()
    if len(n) < 4: return "Name too short"
    if re.search(r"\d", n): return "Name contains digits"
    if not re.match(r"^[A-Za-z][A-Za-z\s\-\.']*$", n): return "Name has odd characters"
    if len(n.split()) < 2: return "Only one word"
    return ""

@st.cache_resource(show_spinner=False)
def get_engine_status():
    from ocr.pipeline import get_engine_status as _status
    return _status()

def _clean_passport_name(name: str) -> str:
    if not name: return ""
    cleaned = name.upper()
    for tok in DEFAULT_NAME_TOKENS: cleaned = cleaned.replace(tok, "")
    return re.sub(r"\s+", " ", cleaned).strip()

def _normalise_date_ethiopian(value: str) -> str:
    if not value: return ""
    value = value.upper().strip()
    month_map = {"JAN":"01","FEB":"02","MAR":"03","APR":"04","MAY":"05","JUN":"06","JUL":"07","AUG":"08","SEP":"09","OCT":"10","NOV":"11","DEC":"12","JANUARY":"01","FEBRUARY":"02","MARCH":"03","APRIL":"04","JUNE":"06","JULY":"07","AUGUST":"08","SEPTEMBER":"09","OCTOBER":"10","NOVEMBER":"11","DECEMBER":"12"}
    m = re.search(r"(\d{1,2})\s*([A-Z]{3,})\s*(\d{2,4})", value)
    if m:
        day, mon, year = m.groups(); day = f"{int(day):02d}"; mon = month_map.get(mon[:3], "01")
        if len(year) == 2:
            yi = int(year); year = f"19{year}" if yi >= 80 else f"20{year}"
        try: datetime(int(year), int(mon), int(day)); return f"{day}/{mon}/{year}"
        except ValueError: pass
    m = re.search(r"(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})", value)
    if m:
        day, mon, year = m.groups()
        if len(year) == 2:
            yi = int(year); year = f"19{year}" if yi >= 80 else f"20{year}"
        try: datetime(int(year), int(mon), int(day)); return f"{int(day):02d}/{int(mon):02d}/{year}"
        except ValueError: pass
    m = re.search(r"(\d{4})[/-](\d{1,2})[/-](\d{1,2})", value)
    if m:
        year, mon, day = m.groups()
        try: datetime(int(year), int(mon), int(day)); return f"{int(day):02d}/{int(mon):02d}/{year}"
        except ValueError: pass
    return value

def _parse_ethiopian_passport(text: str) -> dict:
    text = text.replace("\x00", " ").replace("\r", "\n")
    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
    upper = text.upper()
    out = {"NAME":"","PASSPORT_NO":"","DOB":"","ISSUE_DATE":"","EXPIRY_DATE":"","HOME_ADDRESS":"","MRZ_RAW":""}
    mrz_lines = []
    for line in lines:
        clean = re.sub(r"[^A-Z0-9<]", "", line.upper())
        if 28 <= len(clean) <= 50 and re.search(r"[A-Z0-9<]{25,}", clean): mrz_lines.append(clean)
    if len(mrz_lines) >= 2:
        out["MRZ_RAW"] = mrz_lines[-2] + "\n" + mrz_lines[-1]; mrz1, mrz2 = mrz_lines[-2], mrz_lines[-1]
    else: mrz1, mrz2 = "", ""
    m = re.search(r"\b(E[A-Z]?\d{6,8})\b", upper)
    if m: out["PASSPORT_NO"] = m.group(1)
    elif mrz2:
        m = re.match(r"^([A-Z0-9]{9})", mrz2)
        if m: out["PASSPORT_NO"] = m.group(1)
    if mrz1 and "<<" in mrz1:
        parts = mrz1.split("<<"); surname = parts[0].replace("<", " ").strip()
        if len(surname) > 5 and surname[1] == "<": surname = surname[5:].strip()
        given = parts[1].replace("<", " ").strip(); given = re.sub(r"<+$", "", given).strip()
        name_part = f"{given} {surname}".strip(); name_part = re.sub(r"\s+", " ", name_part)
        if name_part: out["NAME"] = name_part
    else:
        for line in lines:
            if "NAME" in line.upper() and ":" in line:
                parts = line.split(":", 1)
                if len(parts) == 2 and parts[1].strip(): out["NAME"] = parts[1].strip(); break
        if not out["NAME"]:
            for line in lines:
                if re.match(r"^[A-Z]{3,}(?:\s+[A-Z]{2,}){1,3}$", line): out["NAME"] = line; break
    out["NAME"] = _clean_passport_name(out["NAME"])
    if mrz2 and len(mrz2) >= 15:
        m = re.search(r"[A-Z]{3}(\d{6})", mrz2)
        if m:
            yy = int(m.group(1)[0:2]); mm = int(m.group(1)[2:4]); dd = int(m.group(1)[4:6])
            year = 1900 + yy if yy > 30 else 2000 + yy
            try: datetime(year, mm, dd); out["DOB"] = f"{dd:02d}/{mm:02d}/{year}"
            except ValueError: pass
    if not out["DOB"]:
        for line in lines:
            if "BIRTH" in line.upper() or "DOB" in line.upper() or "DATE OF BIRTH" in line.upper():
                m = re.search(r"(\d{1,2}\s*[A-Z]{3,}\s*\d{2,4}|\d{1,2}[/-]\d{1,2}[/-]\d{2,4})", line.upper())
                if m: out["DOB"] = _normalise_date_ethiopian(m.group(1)); break
        if not out["DOB"]:
            all_dates = re.findall(r"\b(\d{1,2}\s*[A-Z]{3,}\s*\d{2,4}|\d{1,2}[/-]\d{1,2}[/-]\d{2,4})\b", upper)
            for d in all_dates:
                norm = _normalise_date_ethiopian(d)
                if norm and re.match(r"\d{2}/\d{2}/\d{4}", norm):
                    parts = norm.split("/"); yr = int(parts[2])
                    if 1940 <= yr <= 2010: out["DOB"] = norm; break
    issue_str = ""; expiry_str = ""
    if mrz2 and len(mrz2) >= 27:
        m = re.search(r"[FM](\d{6})", mrz2)
        if m:
            yy = int(m.group(1)[0:2]); mm = int(m.group(1)[2:4]); dd = int(m.group(1)[4:6])
            year = 1900 + yy if yy > 30 else 2000 + yy
            try: datetime(year, mm, dd); expiry_str = f"{dd:02d}/{mm:02d}/{year}"
            except ValueError: pass
    for line in lines:
        lu = line.upper()
        if ("ISSUE" in lu or "ISSUED" in lu) and not issue_str:
            m = re.search(r"(\d{1,2}\s*[A-Z]{3,}\s*\d{2,4}|\d{1,2}[/-]\d{1,2}[/-]\d{2,4})", lu)
            if m: issue_str = _normalise_date_ethiopian(m.group(1))
        if ("EXPIR" in lu or "VALID" in lu or "DATE OF EXPIRY" in lu) and not expiry_str:
            m = re.search(r"(\d{1,2}\s*[A-Z]{3,}\s*\d{2,4}|\d{1,2}[/-]\d{1,2}[/-]\d{2,4})", lu)
            if m: expiry_str = _normalise_date_ethiopian(m.group(1))
    if not issue_str or not expiry_str:
        all_dates = re.findall(r"\b(\d{1,2}\s*[A-Z]{3,}\s*\d{2,4}|\d{1,2}[/-]\d{1,2}[/-]\d{2,4})\b", upper)
        parsed_dates = []
        for d in all_dates:
            norm = _normalise_date_ethiopian(d)
            if norm and re.match(r"\d{2}/\d{2}/\d{4}", norm): parsed_dates.append(norm)
        if len(parsed_dates) >= 2:
            if not issue_str: issue_str = parsed_dates[0]
            if not expiry_str: expiry_str = parsed_dates[1]
        elif len(parsed_dates) == 1:
            if not expiry_str: expiry_str = parsed_dates[0]
    out["ISSUE_DATE"] = issue_str; out["EXPIRY_DATE"] = expiry_str
    for line in lines:
        lu = line.upper()
        if "PLACE OF BIRTH" in lu or "BIRTH PLACE" in lu or "PLACE" in lu:
            parts = line.split(":", 1)
            if len(parts) == 2 and parts[1].strip(): out["HOME_ADDRESS"] = parts[1].strip(); break
            idx = lu.find("PLACE")
            if idx != -1:
                rem = line[idx+5:].strip()
                if rem: out["HOME_ADDRESS"] = rem; break
    if not out["HOME_ADDRESS"]:
        places = ["ADDIS ABABA","DEBRE MARKOS","DEBRIMARKOS","LIMU","GENET","CHAFETA","JIMMA","JINMA","GONDAR","BAHIR DAR","HAWASSA","MEKELLE","DIRE DAWA","ADAMA","NAZRETH"]
        for line in lines:
            for place in places:
                if place in line.upper(): out["HOME_ADDRESS"] = line.strip(); break
            if out["HOME_ADDRESS"]: break
    if out["DOB"]:
        try:
            d = datetime.strptime(out["DOB"], "%d/%m/%Y"); now = datetime.today()
            out["AGE"] = str(now.year - d.year - ((now.month, now.day) < (d.month, d.day)))
        except Exception: out["AGE"] = ""
    return out

def suggest_marital_children(age_value) -> tuple:
    try: a = int(str(age_value).strip())
    except (ValueError, TypeError): return "SINGLE", "NIL"
    if a <= 25: return "SINGLE", "NIL"
    if a <= 27: return "MARRIED", "1"
    if a <= 31: return "MARRIED", "2"
    if a <= 35: return "MARRIED", "3"
    return "MARRIED", "4"

def _age_from_dob(dob: str) -> str:
    try:
        d = datetime.strptime(dob, "%d/%m/%Y"); now = datetime.today()
        return str(now.year - d.year - ((now.month, now.day) < (d.month, d.day)))
    except Exception: return ""

def _subtract_passport_years(expiry: str, years: int) -> str:
    try:
        value = datetime.strptime(expiry, "%d/%m/%Y")
        try: issue = value.replace(year=value.year - years) + timedelta(days=1)
        except ValueError: issue = value.replace(year=value.year - years, month=2, day=28) + timedelta(days=1)
        return issue.strftime("%d/%m/%Y")
    except (TypeError, ValueError): return ""

def _add_passport_years(issue: str, years: int) -> str:
    try:
        value = datetime.strptime(issue, "%d/%m/%Y")
        try: expiry = value.replace(year=value.year + years) - timedelta(days=1)
        except ValueError: expiry = value.replace(year=value.year + years, month=2, day=28) - timedelta(days=1)
        return expiry.strftime("%d/%m/%Y")
    except (TypeError, ValueError): return ""

def _date_sanity_warnings(dob: str, issue: str, expiry: str) -> list:
    """Return list of human-readable warnings for obviously bad dates."""
    warns = []
    today = datetime.today()
    if dob:
        try:
            d = datetime.strptime(dob, "%d/%m/%Y")
            age = (today - d).days / 365.25
            if age < 18: warns.append(f"Age looks too young ({int(age)} yrs) — check DOB.")
            if age > 60: warns.append(f"Age looks too old ({int(age)} yrs) — check DOB.")
            if d > today: warns.append("Date of birth is in the future.")
        except Exception:
            warns.append("Date of birth is not in DD/MM/YYYY format.")
    if issue:
        try:
            i = datetime.strptime(issue, "%d/%m/%Y")
            if i > today: warns.append("Issue date is in the future.")
            if i.year < 2000: warns.append("Issue date is before 2000 — check.")
        except Exception:
            warns.append("Issue date is not in DD/MM/YYYY format.")
    if expiry:
        try:
            e = datetime.strptime(expiry, "%d/%m/%Y")
            if e < today: warns.append("Passport has already expired.")
        except Exception:
            warns.append("Expiry date is not in DD/MM/YYYY format.")
    if issue and expiry:
        try:
            i = datetime.strptime(issue, "%d/%m/%Y")
            e = datetime.strptime(expiry, "%d/%m/%Y")
            if e <= i: warns.append("Expiry date is not after issue date.")
            gap = (e - i).days
            if gap < 365: warns.append(f"Only {gap} days between issue and expiry — check.")
        except Exception: pass
    return warns

def _merge_info(base: dict, extra: dict) -> dict:
    for k in ["NAME","PASSPORT_NO","DOB","ISSUE_DATE","EXPIRY_DATE","HOME_ADDRESS"]:
        if not base.get(k) and extra.get(k): base[k] = extra[k]
    if base.get("NAME"): base["NAME"] = _clean_passport_name(base["NAME"])
    if base.get("DOB"): base["AGE"] = _age_from_dob(base["DOB"])
    return base

@st.cache_data(show_spinner=False)
def cached_smart_read(passport_bytes, ocr_bytes, use_easyocr, use_gemini,
                       gemini_keys_text, gemini_model, amharic, fast_mode, validity_years=5):
    passport_img = Image.open(BytesIO(passport_bytes)).convert("RGB")
    ocr_img = Image.open(BytesIO(ocr_bytes)).convert("RGB") if ocr_bytes else None
    return run_smart_read(passport_img, ocr_img, use_easyocr, use_gemini,
                          gemini_keys_text, gemini_model, amharic, fast_mode, validity_years=validity_years)

def run_smart_read(passport_img, ocr_img, use_easyocr=True, use_gemini=False,
                    gemini_keys_text=" ", gemini_model="gemini-3.7-flash",
                    amharic=False, fast_mode=True, validity_years=5):
    info = {"NAME":"","PASSPORT_NO":"","DOB":"","EXPIRY_DATE":"","ISSUE_DATE":"","HOME_ADDRESS":"","AGE":"",
            "_warning":"","_mrz":{},"_ocr_text":"","_engine":"local-fast"}
    try:
        from ocr.pipeline import read_passport_local
        local = read_passport_local(passport_img, ocr_image=ocr_img, use_easyocr=use_easyocr)
        ocr_text = local.get("_ocr_text", "") or ""
        info["_engine"] = local.get("_engine", "local"); info["_mrz"] = local.get("_mrz", {}) or {}
        parsed = _parse_ethiopian_passport(ocr_text); _merge_info(info, parsed); info["_ocr_text"] = ocr_text
    except Exception:
        try:
            from ocr.pipeline import _easyocr_text, _ocr_text
            ocr_text = _easyocr_text(passport_img) if use_easyocr else _ocr_text(passport_img)
            parsed = _parse_ethiopian_passport(ocr_text); _merge_info(info, parsed)
            info["_ocr_text"] = ocr_text; info["_engine"] = "fallback"
        except Exception as e2: info["_warning"] = f"OCR error: {e2}"
    keys = []
    if use_gemini and gemini_keys_text.strip():
        try:
            from ocr.pipeline import gemini_keys_from_text
            keys = gemini_keys_from_text(gemini_keys_text)
        except Exception: keys = []
    missing_before = [k for k in ["NAME","PASSPORT_NO","DOB","EXPIRY_DATE","HOME_ADDRESS"] if not info.get(k)]
    if keys and missing_before:
        try:
            from ocr.pipeline import gemini_json_with_rotation
            prompt = ("You are an expert at reading Ethiopian passports. "
                      "Extract the following fields from the passport bio page image and return ONLY a JSON object: "
                      "NAME (Full name in the format GIVEN-NAME(S) SURNAME as it appears on the passport, "
                      "do NOT include the placeholder tokens PQETH or PPETH), "
                      "PASSPORT_NO, DOB (DD/MM/YYYY), ISSUE_DATE (DD/MM/YYYY), "
                      "EXPIRY_DATE (DD/MM/YYYY), HOME_ADDRESS. "
                      "If a field is not found, use empty string. No markdown.")
            data, _, errors = gemini_json_with_rotation(ocr_img or passport_img, keys, gemini_model, prompt)
            if data: _merge_info(info, data)
            if errors and not data:
                info["_warning"] = (info["_warning"] + " " if info["_warning"] else "") + "Gemini unavailable."
        except Exception:
            info["_warning"] = (info["_warning"] + " " if info["_warning"] else "") + "Gemini unavailable."
    if amharic and keys and any(not info.get(k) for k in ["NAME","DOB","EXPIRY_DATE","HOME_ADDRESS"]):
        try:
            from amharic.ocr import tesseract_amharic
            am_text, _ = tesseract_amharic(ocr_img or passport_img)
            if am_text.strip():
                from ocr.pipeline import gemini_text_with_rotation
                instruction = ("Translate this Amharic passport text to English and return ONLY JSON "
                               "with NAME, PASSPORT_NO, DOB, ISSUE_DATE, EXPIRY_DATE, HOME_ADDRESS. "
                               "Dates must be DD/MM/YYYY. Omit PQETH and PPETH from NAME.")
                resp, _, _ = gemini_text_with_rotation(am_text, keys, gemini_model, instruction)
                m = re.search(r"\{.*\}", resp or "", re.DOTALL)
                if m:
                    _merge_info(info, json.loads(m.group()))
        except Exception: pass
    info["NAME"] = _clean_passport_name(info.get("NAME", ""))
    if info.get("EXPIRY_DATE") and not info.get("ISSUE_DATE"):
        ci = _subtract_passport_years(info["EXPIRY_DATE"], validity_years)
        if ci: info["ISSUE_DATE"] = ci
    elif info.get("ISSUE_DATE") and not info.get("EXPIRY_DATE"):
        ce = _add_passport_years(info["ISSUE_DATE"], validity_years)
        if ce: info["EXPIRY_DATE"] = ce
    missing = [k for k in ["NAME","PASSPORT_NO","DOB","EXPIRY_DATE","HOME_ADDRESS"] if not info.get(k)]
    if missing:
        info["_warning"] = (info["_warning"] + " " if info["_warning"] else "") + "Some fields could not be read automatically."
    return info

def find_template(agency, level):
    folder = BASE_DIR / agency
    for n in TEMPLATE_CANDIDATES[level]:
        p = folder / n
        if p.exists(): return p
    return None

def desktop_root():
    home = Path.home()
    for p in [home / "Desktop", Path(os.environ.get("USERPROFILE", "")) / "Desktop"]:
        if p.exists(): return p
    return Path(tempfile.gettempdir())

def safe_name(name): return re.sub(r'[\/*?:"<>|]', "", name).strip() or "candidate"

def date_parts(s):
    try:
        d = datetime.strptime(s, "%d/%m/%Y"); return d.day, d.month, d.year
    except Exception: return None

def _seed_country(slot_value, choice_key, manual_key):
    if choice_key in st.session_state: return
    sv = (slot_value or "").strip()
    if not sv:
        st.session_state[choice_key] = COUNTRY_PRESETS[0]; st.session_state[manual_key] = ""; return
    for p in COUNTRY_PRESETS:
        if p.upper() == sv.upper():
            st.session_state[choice_key] = p; st.session_state[manual_key] = ""; return
    st.session_state[choice_key] = OTHER_LABEL; st.session_state[manual_key] = sv

def _date_keys(label: str):
    base = re.sub(r"[^A-Za-z0-9_]", "_", label)
    return base + "_day", base + "_month", base + "_year"

def _seed_date_widgets(label: str, value: str, min_year: int, max_year: int):
    dk, mk, yk = _date_keys(label)
    parsed = date_parts(value)
    if parsed and min_year <= parsed[2] <= max_year: d, m, y = parsed
    else: d = m = y = 0
    st.session_state[dk] = d; st.session_state[mk] = m; st.session_state[yk] = y

def date_picker(label, min_year, max_year):
    dk, mk, yk = _date_keys(label)
    st.session_state.setdefault(dk, 0); st.session_state.setdefault(mk, 0); st.session_state.setdefault(yk, 0)
    days = [0] + list(range(1, 32)); years = [0] + list(range(min_year, max_year + 1))
    c1, c2, c3 = st.columns([1, 1.35, 1], gap="small")
    d = c1.selectbox(f"{label} day", days, format_func=lambda x: "Day" if x == 0 else str(x), key=dk, label_visibility="collapsed")
    m = c2.selectbox(f"{label} month", list(range(13)), format_func=lambda x: MONTH_OPTIONS[x], key=mk, label_visibility="collapsed")
    y = c3.selectbox(f"{label} year", years, format_func=lambda x: "Year" if x == 0 else str(x), key=yk, label_visibility="collapsed")
    if not d or not m or not y: return ""
    try: datetime(y, m, d); return f"{d:02d}/{m:02d}/{y}"
    except Exception: return ""

def _fit_image_to_box(img, target_w_in, target_h_in, mode="contain", dpi=300):
    tw = max(1, int(round(target_w_in * dpi))); th = max(1, int(round(target_h_in * dpi)))
    W, H = img.size
    if W <= 0 or H <= 0: return img, target_w_in, target_h_in
    if mode == "cover":
        scale = max(tw / W, th / H); nW = max(1, int(round(W * scale))); nH = max(1, int(round(H * scale)))
        scaled = img.resize((nW, nH), Image.Resampling.LANCZOS)
        left = max(0, (nW - tw) // 2); top = max(0, (nH - th) // 2)
        return scaled.crop((left, top, left + tw, top + th)), target_w_in, target_h_in
    scale = min(tw / W, th / H); nW = max(1, int(round(W * scale))); nH = max(1, int(round(H * scale)))
    scaled = img.resize((nW, nH), Image.Resampling.LANCZOS)
    return scaled, nW / dpi, nH / dpi

def _iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables: yield from _iter_table_paragraphs(nested)

def _iter_all_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables: yield from _iter_table_paragraphs(table)

def _clear_runs(paragraph):
    for run in list(paragraph.runs): run._element.getparent().remove(run._element)

def _lock_containing_cell(paragraph, w_in, h_in):
    tc = paragraph._p.getparent()
    while tc is not None and tc.tag != qn("w:tc"): tc = tc.getparent()
    if tc is None: return
    tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.find(qn('w:tcMar'))
    if tcMar is None: tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m in ('top', 'left', 'bottom', 'right'):
        node = tcMar.find(qn(f'w:{m}'))
        if node is None: node = OxmlElement(f'w:{m}'); tcMar.append(node)
        node.set(qn('w:w'), '0'); node.set(qn('w:type'), 'dxa')
    w_dxa = int(round(w_in * 1440)); h_dxa = int(round(h_in * 1440))
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None: tcW = OxmlElement("w:tcW"); tcPr.append(tcW)
    tcW.set(qn("w:type"), "dxa"); tcW.set(qn("w:w"), str(w_dxa))
    tr = tc.getparent()
    if tr is not None and tr.tag == qn("w:tr"):
        trPr = tr.find(qn("w:trPr"))
        if trPr is None: trPr = OxmlElement("w:trPr"); tr.insert(0, trPr)
        trH = trPr.find(qn("w:trHeight"))
        if trH is None: trH = OxmlElement("w:trHeight"); trPr.append(trH)
        ex = int(trH.get(qn("w:val")) or 0); trH.set(qn("w:hRule"), "exact"); trH.set(qn("w:val"), str(max(h_dxa, ex)))
    tbl = tc.getparent()
    while tbl is not None and tbl.tag != qn("w:tbl"): tbl = tbl.getparent()
    if tbl is not None:
        tblPr = tbl.find(qn("w:tblPr"))
        if tblPr is None: tblPr = OxmlElement("w:tblPr"); tbl.insert(0, tblPr)
        tblL = tblPr.find(qn("w:tblLayout"))
        if tblL is None: tblL = OxmlElement("w:tblLayout"); tblPr.append(tblL)
        tblL.set(qn("w:type"), "fixed")

def _image_box_for_key(key):
    if key == "IMAGE_3X4":
        k = float(st.session_state.get("k_3x4", DEFAULT_K_3X4)); bw, bh = BOX_3X4_IN; return (bw*k, bh*k)
    if key == "IMAGE_FULL":
        k = float(st.session_state.get("k_full", DEFAULT_K_FULL)); bw, bh = BOX_FULL_IN; return (bw*k, bh*k)
    if key == "IMAGE_PASSPORT":
        k = float(st.session_state.get("k_passport", DEFAULT_K_PASSPORT)); bw, bh = BOX_PASSPORT_BASE_IN; return (bw*k, bh*k)
    return (2.0, 2.6)

def _process_paragraph(paragraph, data, images):
    full_text = paragraph.text
    if not full_text or "{{" not in full_text: return
    for key, path in images.items():
        ph = "{{%s}}" % key
        if ph in full_text and path:
            _clear_runs(paragraph); run = paragraph.add_run()
            w_in, h_in = _image_box_for_key(key); mode = FIT_MODE.get(key, "contain")
            source = Image.open(path).convert("RGB")
            fitted, aw, ah = _fit_image_to_box(source, w_in, h_in, mode=mode)
            ib = BytesIO(); fitted.save(ib, format="JPEG", quality=95); ib.seek(0)
            run.add_picture(ib, width=Inches(aw), height=Inches(ah))
            if key in ("IMAGE_PASSPORT","IMAGE_3X4","IMAGE_FULL"):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = 0; paragraph.paragraph_format.space_after = 0
                paragraph.paragraph_format.line_spacing = Pt(1)
                paragraph.paragraph_format.left_indent = Inches(IMAGE_RIGHT_SHIFT_IN)
                paragraph.paragraph_format.right_indent = Inches(0)
                pPr = paragraph._p.get_or_add_pPr(); sp = pPr.find(qn('w:spacing'))
                if sp is None: sp = OxmlElement('w:spacing'); pPr.append(sp)
                sp.set(qn('w:before'), '0'); sp.set(qn('w:after'), '0')
                sp.set(qn('w:line'), '240'); sp.set(qn('w:lineRule'), 'auto')
                tc = paragraph._p.getparent()
                while tc is not None and tc.tag != qn("w:tc"): tc = tc.getparent()
                if tc is not None:
                    tcPr = tc.get_or_add_tcPr(); vA = tcPr.find(qn("w:vAlign"))
                    if vA is None: vA = OxmlElement("w:vAlign"); tcPr.append(vA)
                    vA.set(qn("w:val"), "center")
            _lock_containing_cell(paragraph, w_in, h_in)
            return
    new_text = full_text; used = []
    for key, val in data.items():
        ph = "{{%s}}" % key
        if ph in new_text: new_text = new_text.replace(ph, str(val)); used.append(key)
    if used:
        first = paragraph.runs[0] if paragraph.runs else None
        _clear_runs(paragraph); nr = paragraph.add_run(new_text)
        if first is not None and first._element.rPr is not None:
            nr._element.insert(0, copy.deepcopy(first._element.rPr))
        if any(k in BOLD_RED_KEYS for k in used):
            nr.font.bold = True; nr.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

def _force_calibri(doc):
    for p in _iter_all_paragraphs(doc):
        for run in p.runs:
            run.font.name = "Calibri"; rPr = run._element.get_or_add_rPr()
            rF = rPr.find(qn("w:rFonts"))
            if rF is None: rF = OxmlElement("w:rFonts"); rPr.append(rF)
            for a in ("w:ascii","w:hAnsi","w:eastAsia","w:cs"): rF.set(qn(a), "Calibri")

def _find_soffice():
    for n in ("soffice","libreoffice","soffice.exe"):
        p = shutil.which(n)
        if p: return p
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
              str(Path(os.environ.get("LOCALAPPDATA",""))/"Programs"/"LibreOffice"/"program"/"soffice.exe")):
        if os.path.exists(p): return p
    return None

def fill_cv(template_path, data, images, docx_out, pdf_out):
    doc = Document(str(template_path))
    for p in list(_iter_all_paragraphs(doc)): _process_paragraph(p, data, images)
    _force_calibri(doc); doc.save(str(docx_out))

    # Compute the exact same box sizes (in inches) used for the .docx images,
    # so the PDF renders photos at IDENTICAL dimensions -- not
    # independently-guessed CSS sizes that can drift out of sync.
    image_sizes_in = {}
    for key, path in images.items():
        if not path:
            continue
        try:
            w_in, h_in = _image_box_for_key(key)
            source = Image.open(path).convert("RGB")
            _, aw, ah = _fit_image_to_box(source, w_in, h_in, mode=FIT_MODE.get(key, "contain"))
            image_sizes_in[key] = (aw, ah)
        except Exception:
            pass

    # --- PDF: render from HTML/CSS instead of converting the .docx ---
    # Why not just convert the .docx? Two independent bugs showed up doing
    # that on the server (Linux + LibreOffice), neither of which happens
    # in Microsoft Word on a PC:
    #   1) Arabic text-shaping bugs (fixed at the template level separately)
    #   2) Floating image anchors landing in the wrong spot / overlapping
    #      text -- Word and LibreOffice resolve ambiguous floating-image
    #      anchors differently, and there's no reliable way to force them
    #      to agree from inside the .docx.
    # Rendering from HTML avoids both classes of bug entirely, since every
    # element is placed explicitly rather than "floated". The .docx above
    # is completely unaffected and stays fully editable/perfect in Word.
    try:
        from cv_core.html_pdf import render_pdf_via_html
        if render_pdf_via_html(data, images, pdf_out, image_sizes_in=image_sizes_in):
            return True
    except Exception as e:
        st.warning(f"HTML PDF renderer note: {e}")

    # --- Fallbacks, only used if the HTML renderer above isn't available ---
    soffice = _find_soffice()
    if soffice:
        try:
            with tempfile.TemporaryDirectory() as td:
                subprocess.run([soffice,"--headless","--norestore","--convert-to","pdf","--outdir",td,str(docx_out)], capture_output=True, timeout=90)
                src = Path(td)/(docx_out.stem + ".pdf")
                if src.exists() and src.stat().st_size > 0: shutil.copyfile(src, pdf_out); return True
        except Exception: pass
    try:
        from docx2pdf import convert as docx_to_pdf
        import pythoncom
        if os.name == "nt": pythoncom.CoInitialize()
        docx_to_pdf(str(docx_out), str(pdf_out))
        return pdf_out.exists() and pdf_out.stat().st_size > 0
    except Exception as e:
        st.warning(f"PDF conversion note: {e}"); return False

def crop_with_numbers(img, key, current_box, w, h):
    l, t, r, b = current_box
    c1, c2 = st.columns(2)
    nl = c1.number_input("Left", 0, w, l, step=5, key=f"{key}_left")
    nr = c2.number_input("Right", 0, w, r, step=5, key=f"{key}_right")
    nt = c1.number_input("Top", 0, h, t, step=5, key=f"{key}_top")
    nb = c2.number_input("Bottom", 0, h, b, step=5, key=f"{key}_bottom")
    return (min(nl,nr), min(nt,nb), max(nl,nr), max(nt,nb))

def crop_with_canvas(img, key, current_box, canvas_width=CROP_PREVIEW_WIDTH):
    w, h = img.size; version = st.session_state.get("canvas_version", 0)
    ah = int(round(canvas_width * h / max(1, w)))
    nb = drag_crop_box(img, current_box, color=CROP_COLORS[key], key=f"drag_crop_{key}_{version}",
                       max_width=canvas_width, max_height=ah)
    return clamp_box(nb, w, h)

# ======================================================================
# Page setup
# ======================================================================
st.set_page_config(page_title="CV Automation", page_icon="📄", layout="wide")

if "theme_mode" not in st.session_state: st.session_state.theme_mode = "Gray"
if "k_3x4" not in st.session_state: st.session_state.k_3x4 = DEFAULT_K_3X4
if "k_full" not in st.session_state: st.session_state.k_full = DEFAULT_K_FULL
if "k_passport" not in st.session_state: st.session_state.k_passport = DEFAULT_K_PASSPORT
if "enable_shortcuts" not in st.session_state: st.session_state.enable_shortcuts = True

def _theme_palette(mode):
    if mode == "Dark":
        return dict(bg="#08060f", panel="#0f0a1c", panel2="#06040d", border="#241a3a",
                    text="#c8b0e8", text_soft="#7a6a95", btn="#120b22", btn_hover="#1c1230",
                    btn_primary="#a78bfa", btn_primary_hover="#c4b5fd", tab_bg="#120b22",
                    input_bg="#0a0718", accent="#f0b860")
    if mode == "Light":
        return dict(bg="#f4f5f7", panel="#ffffff", panel2="#edf0f4", border="#d2d5dc",
                    text="#111827", text_soft="#4b5563", btn="#374151", btn_hover="#1f2937",
                    btn_primary="#2563eb", btn_primary_hover="#1d4ed8", tab_bg="#e5e7eb",
                    input_bg="#ffffff", accent="#2563eb")
    if mode == "Colorful":
        return dict(bg="#faf7f0", panel="#fffdf7", panel2="#f3ede1", border="#e8dfd1",
                    text="#2d2a24", text_soft="#8a8070", btn="#d97757", btn_hover="#c66648",
                    btn_primary="#d97757", btn_primary_hover="#c66648", tab_bg="#f0e8d8",
                    input_bg="#fffdf7", accent="#cc785c")
    return dict(bg="#dcdee2", panel="#ebedf0", panel2="#d5d8dd", border="#bdc1c9",
                text="#16171a", text_soft="#3c4046", btn="#475569", btn_hover="#334155",
                btn_primary="#4b5563", btn_primary_hover="#374151", tab_bg="#ced1d7",
                input_bg="#f4f5f7", accent="#334155")

def _theme_css(mode):
    t = _theme_palette(mode)
    bg, panel, panel2, border = t["bg"], t["panel"], t["panel2"], t["border"]
    text, text_soft = t["text"], t["text_soft"]
    btn, btn_hover = t["btn"], t["btn_hover"]
    btn_primary, btn_primary_hover = t["btn_primary"], t["btn_primary_hover"]
    tab_bg, input_bg, accent = t["tab_bg"], t["input_bg"], t["accent"]
    is_dark = (mode == "Dark"); is_colorful = (mode == "Colorful")
    if is_dark:
        btn_text_color = "#c4b5fd"; border_radius = "4px"
        hls = "0.14em"; ltf = "uppercase"; lls = "0.14em"
        sidebar_border = "1px solid #241a3a"; primary_text = "#08060f"
    elif is_colorful:
        btn_text_color = "#8a4a35"; border_radius = "10px"
        hls = "0.04em"; ltf = "none"; lls = "0"
        sidebar_border = "1px solid #e8dfd1"; primary_text = "#fffdf7"
    else:
        btn_text_color = "#ffffff"; border_radius = "4px"
        hls = "0"; ltf = "none"; lls = "0"
        sidebar_border = f"1px solid {border}"; primary_text = "#ffffff"

    dark_extras = ""
    if is_dark:
        dark_extras = f"""
        .stApp {{ background: #08060f !important; background-image: radial-gradient(circle at 20% 10%, rgba(167,139,250,0.10) 0%, transparent 50%), radial-gradient(circle at 80% 90%, rgba(96,165,250,0.08) 0%, transparent 50%), radial-gradient(circle at 50% 50%, rgba(240,184,96,0.04) 0%, transparent 60%) !important; }}
        h1, h2, h3, h4 {{ text-transform: uppercase; letter-spacing: {hls}; font-weight: 400 !important; color: #c8b0e8 !important; -webkit-text-fill-color: #c8b0e8; }}
        .hero-title {{ text-transform: uppercase; letter-spacing: 0.4em; font-weight: 200 !important; font-size: 1.5rem !important; background: linear-gradient(135deg, #c8b0e8 0%, #a78bfa 40%, #60a5fa 70%, #f0b860 100%) !important; -webkit-background-clip: text !important; -webkit-text-fill-color: transparent !important; background-clip: text !important; border-bottom: 1px solid #241a3a; padding-bottom: 14px; margin-bottom: 8px !important; }}
        .hero-sub {{ letter-spacing: 0.18em; text-transform: uppercase; font-size: 0.65rem !important; color: #60a5fa !important; -webkit-text-fill-color: #60a5fa; }}
        div[data-testid="stTextInput"] label, div[data-testid="stSelectbox"] label {{ text-transform: uppercase; letter-spacing: 0.14em !important; font-weight: 400 !important; color: #9a8fb8 !important; -webkit-text-fill-color: #9a8fb8; font-size: 0.7rem !important; }}
        .details-header, .date-section-title, .config-title, .live-crop-label, .ocr-header {{ text-transform: uppercase; letter-spacing: 0.18em !important; font-weight: 400 !important; color: #a78bfa !important; -webkit-text-fill-color: #a78bfa; border-bottom: none !important; }}
        .k-row-label {{ text-transform: uppercase; letter-spacing: 0.1em !important; font-weight: 400 !important; color: #f0b860 !important; -webkit-text-fill-color: #f0b860; }}
        div[data-testid="stVerticalBlockBorderWrapper"] {{ border-radius: 4px !important; box-shadow: none !important; border: 1px solid #241a3a !important; background: #0f0a1c !important; }}
        .compact-card {{ border-radius: 4px !important; border: none !important; background: transparent !important; }}
        .config-bar {{ border-radius: 4px !important; border-left: 1px solid #a78bfa !important; background: #0f0a1c !important; }}
        input, textarea, .stSelectbox div[data-baseweb="select"] > div {{ border-radius: 4px !important; border: 1px solid #241a3a !important; background-color: #0a0718 !important; color: #c8b0e8 !important; }}
        input:focus, textarea:focus {{ border-color: #a78bfa !important; box-shadow: 0 0 0 2px rgba(167,139,250,0.25) !important; }}
        .stButton > button {{ border-radius: 4px !important; border: 1px solid #a78bfa !important; background: transparent !important; color: #c4b5fd !important; text-transform: uppercase; letter-spacing: 0.22em !important; font-weight: 400 !important; font-size: 0.65rem !important; transition: all 0.25s ease; }}
        .stButton > button:hover {{ background: linear-gradient(135deg, #a78bfa 0%, #60a5fa 100%) !important; color: #08060f !important; border-color: #a78bfa !important; box-shadow: 0 0 20px rgba(167,139,250,0.45) !important; }}
        .stButton > button[kind="primary"] {{ background: linear-gradient(135deg, #a78bfa 0%, #60a5fa 100%) !important; color: #08060f !important; border: 1px solid #a78bfa !important; }}
        .stButton > button[kind="primary"]:hover {{ background: linear-gradient(135deg, #c4b5fd 0%, #93c5fd 100%) !important; border-color: #c4b5fd !important; color: #08060f !important; box-shadow: 0 0 24px rgba(167,139,250,0.55) !important; }}
        .mode-pill div[data-testid="stRadio"] > div[role="radiogroup"] > label {{ border-radius: 4px !important; background: transparent !important; border: 1px solid #241a3a !important; color: #7a6a95 !important; text-transform: uppercase; letter-spacing: 0.14em; font-size: 0.65rem !important; }}
        .mode-pill div[data-testid="stRadio"] > div[role="radiogroup"] > label:has(input:checked) {{ background: #a78bfa !important; border-color: #a78bfa !important; color: #08060f !important; }}
        .mode-pill div[data-testid="stRadio"] > div[role="radiogroup"] > label:has(input:checked) p,
        .mode-pill div[data-testid="stRadio"] > div[role="radiogroup"] > label:has(input:checked) span {{ color: #08060f !important; }}
        section[data-testid="stSidebar"] {{ border-right: 1px solid #241a3a !important; background: #06040d !important; }}
        .k-row div[data-testid="stRadio"] > div[role="radiogroup"] > label {{ border-radius: 4px !important; border: 1px solid #241a3a !important; background: transparent !important; color: #7a6a95 !important; text-transform: uppercase; }}
        .k-row div[data-testid="stRadio"] > div[role="radiogroup"] > label:has(input:checked) {{ background: #f0b860 !important; border-color: #f0b860 !important; color: #08060f !important; }}
        .k-row div[data-testid="stRadio"] > div[role="radiogroup"] > label:has(input:checked) p,
        .k-row div[data-testid="stRadio"] > div[role="radiogroup"] > label:has(input:checked) span {{ color: #08060f !important; }}
        .step-tag {{ background: linear-gradient(135deg, #a78bfa 0%, #60a5fa 100%) !important; border: none !important; color: #08060f !important; border-radius: 999px !important; text-transform: uppercase; letter-spacing: 0.25em; font-weight: 600 !important; }}
        .derived-date {{ border-radius: 4px !important; border: 1px solid #241a3a !important; border-left: 2px solid #f0b860 !important; background: #0f0a1c !important; }}
        .derived-date .dd-value {{ color: #f0b860 !important; background: none !important; -webkit-text-fill-color: #f0b860 !important; }}
        .derived-date .dd-label {{ letter-spacing: 0.25em; color: #60a5fa !important; font-weight: 400 !important; }}
        .derived-date .dd-hint {{ color: #7a6a95 !important; }}
        details[data-testid="stExpander"] {{ border-radius: 4px !important; border: 1px solid #241a3a !important; background: #0f0a1c !important; }}
        div[data-baseweb="popover"], div[data-baseweb="menu"], ul[role="listbox"] {{ border-radius: 4px !important; border: 1px solid #241a3a !important; background: #0f0a1c !important; }}
        li[role="option"]:hover, li[role="option"][aria-selected="true"] {{ background: #1c1230 !important; color: #c8b0e8 !important; }}
        section[data-testid="stFileUploadDropzone"] {{ border-radius: 4px !important; border: 1px dashed #241a3a !important; background: #0a0718 !important; }}
        div[data-testid="stProgressBar"] > div > div > div {{ background: linear-gradient(90deg, #a78bfa 0%, #60a5fa 50%, #f0b860 100%) !important; }}
        div[data-testid="stProgressBar"] > div {{ background-color: #120b22 !important; }}
        .crop-caption {{ background: #0f0a1c !important; border: 1px solid #241a3a !important; color: #9a8fb8 !important; letter-spacing: 0.08em; }}
        .crop-legend {{ color: #60a5fa !important; letter-spacing: 0.1em; }}
        .section-divider {{ border-top: 1px solid #241a3a !important; }}
        .config-title span {{ background: linear-gradient(90deg, #a78bfa, #60a5fa, #f0b860) !important; -webkit-background-clip: text !important; -webkit-text-fill-color: transparent !important; background-clip: text !important; font-weight: 400 !important; }}
        .crop-help-box {{ background: #0f0a1c !important; border: 1px solid #241a3a !important; border-left: 2px solid #a78bfa !important; border-radius: 4px !important; }}
        .crop-help-title {{ color: #a78bfa !important; letter-spacing: 0.15em; font-weight: 400 !important; }}
        .crop-help-body {{ color: #7a6a95 !important; }}
        .crop-help-body b {{ color: #c8b0e8 !important; }}
        .preview-hero {{ background: linear-gradient(135deg, rgba(167,139,250,0.12) 0%, rgba(96,165,250,0.12) 100%) !important; border: 1px solid #241a3a !important; border-radius: 4px !important; }}
        .preview-hero-title {{ background: linear-gradient(90deg, #a78bfa, #60a5fa) !important; -webkit-background-clip: text !important; -webkit-text-fill-color: transparent !important; background-clip: text !important; letter-spacing: 0.15em; font-weight: 400 !important; }}
        .preview-hero-badge {{ background: transparent !important; border: 1px solid #f0b860 !important; color: #f0b860 !important; border-radius: 4px !important; letter-spacing: 0.2em !important; }}
        .nudge-header {{ color: #a78bfa !important; letter-spacing: 0.15em; }}
        .info-bar {{ background: #0f0a1c !important; border-left: 3px solid #60a5fa !important; color: #9a8fb8 !important; }}
        .warn-bar {{ background: #0f0a1c !important; border-left: 3px solid #f0b860 !important; color: #f0b860 !important; }}
        div[data-testid="stAlert"] {{ background: #0f0a1c !important; border: 1px solid #241a3a !important; border-radius: 4px !important; }}
        div[data-testid="stAlert"] p {{ color: #c8b0e8 !important; }}
        div[data-testid="stCaptionContainer"] p {{ color: #60a5fa !important; }}
        ::-webkit-scrollbar {{ width: 8px; height: 8px; }}
        ::-webkit-scrollbar-track {{ background: #08060f; }}
        ::-webkit-scrollbar-thumb {{ background: #241a3a; border-radius: 4px; }}
        ::-webkit-scrollbar-thumb:hover {{ background: #a78bfa; }}
        div[data-baseweb="popover"], div[data-baseweb="popover"] *,
        div[data-baseweb="menu"], div[data-baseweb="menu"] *,
        ul[role="listbox"], ul[role="listbox"] *,
        li[role="option"],
        div[data-baseweb="select"] div[role="listbox"],
        div[data-baseweb="select"] ul {{ background-color: #1a1030 !important; color: #c8b0e8 !important; border-color: #241a3a !important; }}
        li[role="option"]:hover, li[role="option"][aria-selected="true"], li[aria-selected="true"] {{ background: linear-gradient(90deg, #a78bfa 0%, #60a5fa 100%) !important; color: #08060f !important; }}
        li[role="option"] span, li[role="option"] p,
        div[data-baseweb="menu"] span, div[data-baseweb="menu"] p {{ color: #c8b0e8 !important; }}
        li[role="option"]:hover span, li[role="option"]:hover p,
        li[role="option"][aria-selected="true"] span, li[role="option"][aria-selected="true"] p {{ color: #08060f !important; }}
        div[data-baseweb="select"] > div {{ background-color: #0a0718 !important; border-color: #241a3a !important; color: #c8b0e8 !important; }}
        div[data-baseweb="select"] input {{ color: #c8b0e8 !important; background: transparent !important; }}
        div[data-baseweb="select"] svg {{ fill: #a78bfa !important; }}
        """

    colorful_extras = ""
    if is_colorful:
        colorful_extras = f"""
        .stApp {{ background: radial-gradient(circle at 8% 3%, rgba(217,119,87,0.06) 0%, transparent 40%), radial-gradient(circle at 92% 97%, rgba(200,168,124,0.06) 0%, transparent 40%), #faf7f0 !important; }}
        h1, h2, h3, h4 {{ color: #3d3229 !important; letter-spacing: 0.005em; font-weight: 700; }}
        .hero-title {{ color: #3d3229 !important; font-weight: 800 !important; border-bottom: 2px solid #d97757; padding-bottom: 10px; margin-bottom: 6px !important; letter-spacing: 0.02em; }}
        .hero-sub {{ color: #8a8070 !important; font-weight: 500; }}
        div[data-testid="stTextInput"] label, div[data-testid="stSelectbox"] label {{ color: #6b5e4f !important; font-weight: 700 !important; }}
        .details-header {{ color: #cc785c !important; font-weight: 800 !important; border-bottom: none !important; letter-spacing: 0.05em; }}
        .date-section-title {{ color: #d97757 !important; font-weight: 800 !important; }}
        .config-title {{ color: #3d3229 !important; }}
        .config-title span {{ color: #cc785c !important; font-weight: 800 !important; }}
        .k-row-label {{ color: #8a6a4a !important; font-weight: 800 !important; }}
        .live-crop-label {{ color: #3d3229 !important; font-weight: 800 !important; }}
        div[data-testid="stVerticalBlockBorderWrapper"] {{ border: 1px solid #e8dfd1 !important; border-radius: {border_radius} !important; background: #fffdf7 !important; box-shadow: 0 2px 12px rgba(217,119,87,0.06) !important; }}
        .compact-card {{ border: 1px solid #e8dfd1 !important; background: #fffdf7 !important; }}
        .config-bar {{ border-left: 4px solid #d97757 !important; background: #fffdf7 !important; }}
        input, textarea, .stSelectbox div[data-baseweb="select"] > div {{ border: 1.5px solid #e8dfd1 !important; border-radius: {border_radius} !important; background-color: #fffdf7 !important; color: #2d2a24 !important; }}
        input:focus, textarea:focus {{ border-color: #d97757 !important; box-shadow: 0 0 0 3px rgba(217,119,87,0.15) !important; }}
        .stButton > button {{ border-radius: {border_radius} !important; border: 1.5px solid #d97757 !important; background: #fffdf7 !important; color: #cc785c !important; font-weight: 700 !important; transition: all 0.2s ease; }}
        .stButton > button:hover {{ background: #d97757 !important; color: #fffdf7 !important; box-shadow: 0 4px 14px rgba(217,119,87,0.25); }}
        .stButton > button[kind="primary"] {{ background: #d97757 !important; color: #fffdf7 !important; border: 1.5px solid #d97757 !important; box-shadow: 0 3px 12px rgba(217,119,87,0.28) !important; }}
        .stButton > button[kind="primary"]:hover {{ background: #c66648 !important; border-color: #c66648 !important; box-shadow: 0 5px 18px rgba(217,119,87,0.4) !important; }}
        .mode-pill div[data-testid="stRadio"] > div[role="radiogroup"] > label {{ border-radius: {border_radius} !important; border: 1.5px solid #e8dfd1 !important; background: #fffdf7 !important; color: #8a8070 !important; font-weight: 600; }}
        .mode-pill div[data-testid="stRadio"] > div[role="radiogroup"] > label:has(input:checked) {{ background: #d97757 !important; border-color: #d97757 !important; color: #fffdf7 !important; }}
        .mode-pill div[data-testid="stRadio"] > div[role="radiogroup"] > label:has(input:checked) p,
        .mode-pill div[data-testid="stRadio"] > div[role="radiogroup"] > label:has(input:checked) span {{ color: #fffdf7 !important; }}
        section[data-testid="stSidebar"] {{ background: #f3ede1 !important; border-right: 1.5px solid #e8dfd1 !important; }}
        .k-row div[data-testid="stRadio"] > div[role="radiogroup"] > label {{ border-radius: {border_radius} !important; border: 1.5px solid #e8dfd1 !important; background: #fffdf7 !important; color: #8a8070 !important; font-weight: 600; }}
        .k-row div[data-testid="stRadio"] > div[role="radiogroup"] > label:has(input:checked) {{ background: #d97757 !important; border-color: #d97757 !important; color: #fffdf7 !important; }}
        .k-row div[data-testid="stRadio"] > div[role="radiogroup"] > label:has(input:checked) p,
        .k-row div[data-testid="stRadio"] > div[role="radiogroup"] > label:has(input:checked) span {{ color: #fffdf7 !important; font-weight: 700; }}
        .step-tag {{ background: #d97757 !important; color: #fffdf7 !important; border-radius: 8px !important; font-weight: 800 !important; box-shadow: 0 2px 6px rgba(217,119,87,0.2); }}
        .step-tag.step-1 {{ background: #d97757 !important; }}
        .step-tag.step-2 {{ background: #c89860 !important; }}
        .step-tag.step-3 {{ background: #8a9a5b !important; }}
        .derived-date {{ border: 1.5px solid #e8dfd1 !important; border-left: 5px solid #8a9a5b !important; background: #f5f1e4 !important; }}
        .derived-date .dd-value {{ color: #cc785c !important; font-weight: 800 !important; }}
        .derived-date .dd-label {{ color: #8a9a5b !important; font-weight: 800 !important; }}
        section[data-testid="stFileUploadDropzone"] {{ border: 1.5px dashed #d97757 !important; background: #faf3ea !important; }}
        div[data-testid="stProgressBar"] > div > div > div {{ background: linear-gradient(90deg, #d97757 0%, #c89860 50%, #8a9a5b 100%) !important; }}
        div[data-testid="stProgressBar"] > div {{ background-color: #f0e8d8 !important; }}
        .crop-caption {{ background: #f5f1e4 !important; border: 1.5px solid #e8dfd1 !important; color: #6b5e4f !important; font-weight: 600; }}
        .crop-legend {{ color: #cc785c !important; font-weight: 600; }}
        .section-divider {{ border-top: 1.5px solid #e8dfd1 !important; opacity: 1 !important; }}
        """

    return f"""
    <style>
    div.block-container {{ padding-top: 2.2rem !important; padding-bottom: 0.4rem !important; padding-left: 1.0rem !important; padding-right: 1.0rem !important; }}
    header[data-testid="stHeader"] {{ height: 0 !important; background: transparent !important; }}
    div[data-testid="stToolbar"] {{ right: 0.5rem; }}
    .stApp {{ background-color: {bg}; color: {text}; }}
    section[data-testid="stSidebar"] {{ background-color: {panel2}; border-right: {sidebar_border}; }}
    section[data-testid="stSidebar"] .stMarkdown {{ margin-bottom: 0.1rem; }}

    .stApp p {{ margin: 0 !important; padding: 0 !important; }}
    .stApp .stMarkdown p {{ margin-bottom: 0.2rem !important; }}
    .stApp div[data-testid="stElementContainer"] {{ margin: 0 !important; }}
    .stApp div[data-testid="stMarkdownContainer"] {{ margin: 0 !important; }}
    .stApp div[data-testid="stHeading"] {{ margin: 0 !important; padding: 0 !important; }}
    .stApp div[data-testid="stHeading"] h2,
    .stApp div[data-testid="stHeading"] h3 {{ margin: 0 !important; padding: 0 !important; line-height: 1.2 !important; }}
    .stApp .stMarkdown > div > p:empty {{ display: none !important; }}
    .stApp .stVerticalBlock {{ gap: 0.22rem !important; }}
    .stApp .stColumn {{ padding: 0 !important; }}

    h1, h2, h3, h4 {{ color: {text}; font-weight: 700; }}
    h1 {{ font-size: 1.3rem !important; margin: 0 !important; padding: 0 !important; }}
    h2, h3 {{ font-size: 0.95rem !important; margin: 0 !important; padding: 0 !important; }}
    p, span, label, li, div, .stCaption, .stMarkdown {{ color: {text}; }}
    small, .stCaption p {{ color: {text_soft}; opacity: 0.85; }}
    div[data-testid="stVerticalBlock"] {{ gap: 0.22rem !important; }}
    div[data-testid="stHorizontalBlock"] {{ gap: 0.4rem !important; }}
    hr {{ margin: 0.15rem 0 !important; border-color: {border} !important; }}

    div[data-testid="stVerticalBlockBorderWrapper"] {{ background-color: {panel}; border: 1px solid {border}; border-radius: {border_radius}; padding: 0.55rem !important; box-shadow: 0 1px 4px rgba(0,0,0,0.04); }}

    input, textarea, .stSelectbox div[data-baseweb="select"] > div {{ color: {text} !important; background-color: {input_bg} !important; border-radius: {border_radius} !important; border: 1px solid {border} !important; min-height: 1.75rem !important; font-size: 0.82rem !important; padding: 1px 8px !important; box-sizing: border-box !important; }}
    div[data-testid="stSelectbox"] div[data-baseweb="select"] > div {{ border-radius: 8px !important; }}
    .stSelectbox div[data-baseweb="select"] span {{ color: {text} !important; font-size: 0.82rem !important; }}
    .stSelectbox svg {{ fill: {text_soft} !important; }}
    div[data-testid="stTextInput"] label, div[data-testid="stSelectbox"] label {{ font-size: 0.78rem !important; font-weight: 600 !important; margin-bottom: 0.15rem !important; color: {text} !important; text-transform: {ltf}; letter-spacing: {lls}; }}

    div[data-baseweb="popover"], div[data-baseweb="menu"], ul[role="listbox"] {{ background-color: {panel} !important; border: 1px solid {border} !important; border-radius: {border_radius} !important; }}
    li[role="option"] {{ color: {text} !important; background-color: {panel} !important; font-size: 0.78rem !important; padding: 3px 8px !important; }}
    li[role="option"]:hover, li[role="option"][aria-selected="true"] {{ background-color: {panel2} !important; color: {text} !important; }}

    div[data-testid="stRadio"] label, div[data-testid="stCheckbox"] label {{ color: {text} !important; font-size: 0.78rem !important; }}
    div[data-testid="stRadio"] label p, div[data-testid="stCheckbox"] label p {{ color: {text} !important; font-size: 0.78rem !important; }}

    .mode-pill div[data-testid="stRadio"] > div[role="radiogroup"] {{ gap: 0.25rem !important; flex-wrap: nowrap !important; }}
    .mode-pill div[data-testid="stRadio"] > div[role="radiogroup"] > label {{ padding: 2px 12px !important; border-radius: 999px !important; background: {input_bg} !important; border: 1px solid {border} !important; font-size: 0.72rem !important; margin: 0 !important; min-height: 24px; cursor: pointer !important; transition: all 0.15s ease; }}
    .mode-pill div[data-testid="stRadio"] > div[role="radiogroup"] > label:hover {{ background: {panel2} !important; }}
    .mode-pill div[data-testid="stRadio"] > div[role="radiogroup"] > label:has(input:checked) {{ background: {btn_primary} !important; border-color: {btn_primary} !important; }}
    .mode-pill div[data-testid="stRadio"] > div[role="radiogroup"] > label:has(input:checked) p,
    .mode-pill div[data-testid="stRadio"] > div[role="radiogroup"] > label:has(input:checked) span {{ color: {primary_text} !important; font-weight: 600 !important; }}
    .mode-pill div[data-testid="stRadio"] > div[role="radiogroup"] > label > div:first-child {{ display: none !important; }}

    .stButton > button {{ background-color: {btn}; color: {btn_text_color if (is_dark or is_colorful) else '#ffffff'} !important; border: none; border-radius: {border_radius}; font-weight: 600; min-height: 1.75rem !important; height: 1.75rem !important; line-height: 1.75rem !important; font-size: 0.75rem !important; padding: 0 6px !important; display: inline-flex !important; align-items: center !important; justify-content: center !important; transition: background-color 0.15s ease; }}
    .stButton > button p, .stButton > button span, .stButton > button div {{ font-size: 0.75rem !important; line-height: 1 !important; }}
    .stButton > button:hover {{ background-color: {btn_hover}; color: {btn_text_color if (is_dark or is_colorful) else '#ffffff'} !important; }}
    .stButton > button[kind="primary"] {{ background-color: {btn_primary}; color: {primary_text} !important; }}
    .stButton > button[kind="primary"]:hover {{ background-color: {btn_primary_hover}; }}

    div[data-testid="stButton"] button[kind="primary"] {{ padding: 0 10px !important; font-size: 0.68rem !important; letter-spacing: 0.05em !important; min-height: 1.55rem !important; height: 1.55rem !important; line-height: 1.55rem !important; white-space: nowrap; }}

    section[data-testid="stFileUploadDropzone"] {{ background-color: {panel} !important; border: 1px dashed {border} !important; border-radius: {border_radius} !important; padding: 0.4rem !important; }}
    section[data-testid="stFileUploadDropzone"] div, section[data-testid="stFileUploadDropzone"] small {{ color: {text_soft} !important; }}

    details[data-testid="stExpander"] {{ background-color: {panel} !important; border: 1px solid {border} !important; border-radius: {border_radius} !important; }}
    details[data-testid="stExpander"] summary span {{ color: {text} !important; font-size: 0.78rem !important; }}

    .step-tag {{ display: inline-block; background-color: {btn_primary}; color: {primary_text} !important; font-size: 0.58rem; font-weight: 700; letter-spacing: 0.04em; padding: 1px 6px; border-radius: 999px; margin-bottom: 0; line-height: 1.2; }}
    .hero-title {{ text-align:center; margin: 0 !important; padding: 0.25rem 0 0.1rem 0 !important; color: {text}; font-size: 1.4rem; font-weight: 700; line-height: 1.2; }}
    .hero-sub {{ text-align:center; margin: 0 0 0.45rem 0 !important; padding: 0 !important; color: {text_soft}; opacity: 0.9; font-size: 0.78rem; line-height: 1.3; }}
    .ocr-header {{ font-size: 1.0rem !important; font-weight: 700; color: {text}; margin: 0 !important; }}

    .compact-card {{ background: transparent !important; border: none !important; border-radius: 0 !important; padding: 0.15rem 0 !important; box-shadow: none !important; }}
    .details-header {{ font-size: 0.92rem !important; font-weight: 700; color: {text}; margin-top: 0.1rem !important; margin-bottom: 0.25rem !important; border-bottom: none !important; padding: 0 !important; }}
    .details-header:first-child {{ margin-top: 0 !important; }}
    .section-divider {{ margin: 0.3rem 0 0.25rem 0 !important; border: none; border-top: 1px solid {border}; opacity: 0.4; }}

    .date-section-title {{ font-size: 0.8rem !important; font-weight: 700 !important; color: {text} !important; display: block; padding: 2px 0 !important; margin: 0 !important; }}
    .date-block-gap {{ height: 0.4rem; width: 100%; display: block; }}

    .derived-date {{ background: {panel2}; border: 1px solid {border}; border-left: 4px solid {accent}; padding: 5px 12px; border-radius: {border_radius}; display: flex; align-items: center; gap: 12px; margin: 0 !important; min-height: 34px; }}
    .derived-date .dd-label {{ font-size: 0.65rem; color: {text_soft}; text-transform: uppercase; letter-spacing: 0.08em; font-weight: 700; white-space: nowrap; }}
    .derived-date .dd-value {{ font-size: 1.02rem; font-weight: 700; color: {accent}; letter-spacing: 0.03em; font-family: 'Consolas', 'Menlo', monospace; }}
    .derived-date .dd-hint {{ margin-left: auto; font-size: 0.68rem; color: {text_soft}; font-style: italic; white-space: nowrap; }}

    .stSelectbox div[data-baseweb="select"] > div {{ font-size: 0.82rem !important; min-height: 1.75rem !important; }}
    .details-space {{ height: 0.25rem; }}
    .field-row-gap {{ height: 0.5rem; width: 100%; display: block; }}

    div[data-testid="stCaptionContainer"] {{ margin-top: 0 !important; margin-bottom: 0.05rem !important; }}
    div[data-testid="stCaptionContainer"] p {{ font-size: 0.72rem !important; }}

    .live-crop-head {{ padding-top: 0; margin: 0 0 5px 0; }}
    .live-crop-label {{ font-size: 1.05rem !important; font-weight: 800 !important; color: {text} !important; letter-spacing: 0.01em; margin: 0 !important; padding: 0 !important; line-height: 1.25; }}

    .crop-caption {{ margin-top: 0 !important; margin-bottom: 4px !important; font-size: 0.78rem; color: {text_soft}; padding: 4px 8px; background: {panel}; border: 1px solid {border}; border-radius: {border_radius}; }}
    .crop-legend {{ margin-top: 0 !important; margin-bottom: 6px !important; font-size: 0.72rem !important; color: {text_soft}; text-align: center; width: 100%; display: block; }}

    .k-row-label {{ font-size: 0.72rem !important; font-weight: 600 !important; color: {text_soft} !important; margin: 4px 0 1px 0 !important; display: block; }}
    .k-row div[data-testid="stRadio"] > div[role="radiogroup"] {{ gap: 0.15rem !important; flex-wrap: wrap !important; }}
    .k-row div[data-testid="stRadio"] > div[role="radiogroup"] > label {{ padding: 1px 6px !important; font-size: 0.68rem !important; min-height: 20px !important; }}

    .config-bar {{ background: {panel}; border: 1px solid {border}; border-left: 4px solid {accent}; border-radius: {border_radius}; padding: 6px 12px 4px 12px; margin-bottom: 6px; }}
    .config-title {{ font-size: 0.9rem !important; font-weight: 800 !important; color: {text} !important; margin-bottom: 4px !important; display: block; }}

    .st-key-sticky_progress {{ position: sticky !important; top: 0 !important; z-index: 100 !important; background-color: {bg} !important; padding: 6px 12px !important; margin: -8px -12px 4px -12px !important; border-bottom: 1px solid {border} !important; }}
    .st-key-sticky_progress div[data-testid="stProgressBar"] {{ margin: 0 !important; }}

    .st-key-approve_bar {{ position: sticky !important; bottom: 0 !important; z-index: 99 !important; background: {bg} !important; padding: 8px 0 10px 0 !important; border-top: 1px solid {border} !important; }}

    .crop-region-row div[data-testid="stRadio"] > div[role="radiogroup"] {{ gap: 0.35rem !important; flex-wrap: wrap !important; }}
    .crop-region-row div[data-testid="stRadio"] > div[role="radiogroup"] > label {{ padding: 5px 12px !important; border-radius: 999px !important; background: {input_bg} !important; border: 1px solid {border} !important; font-size: 0.78rem !important; font-weight: 600 !important; min-height: 28px !important; cursor: pointer !important; transition: all 0.15s ease; }}
    .crop-region-row div[data-testid="stRadio"] > div[role="radiogroup"] > label:hover {{ background: {panel2} !important; }}
    .crop-region-row div[data-testid="stRadio"] > div[role="radiogroup"] > label:has(input:checked) {{ background: {btn_primary} !important; border-color: {btn_primary} !important; }}
    .crop-region-row div[data-testid="stRadio"] > div[role="radiogroup"] > label:has(input:checked) p,
    .crop-region-row div[data-testid="stRadio"] > div[role="radiogroup"] > label:has(input:checked) span {{ color: {primary_text} !important; font-weight: 700 !important; }}
    .crop-region-row div[data-testid="stRadio"] > div[role="radiogroup"] > label > div:first-child {{ display: none !important; }}
    .crop-region-row {{ margin-bottom: 0.35rem !important; }}

    .crop-help-box {{ background: {panel}; border: 1px dashed {border}; border-left: 4px solid {accent}; border-radius: {border_radius}; padding: 8px 12px; margin-top: 6px !important; }}
    .crop-help-title {{ font-weight: 800; font-size: 0.82rem; margin-bottom: 5px; color: {accent}; letter-spacing: 0.02em; }}
    .crop-help-body {{ font-size: 0.76rem; line-height: 1.5; color: {text_soft}; }}
    .crop-help-body b {{ color: {text}; font-weight: 700; }}

    /* ============ Aurora over mountains · shooting stars ============ */
    .crop-motivation {{
        margin-top: 8px !important;
        padding: 0 !important;
        border-radius: 8px;
        border: 1px solid #241a3a;
        height: 320px;
        min-height: 320px;
        position: relative;
        overflow: hidden;
        background: linear-gradient(180deg, #03020a 0%, #08061a 30%, #100a24 55%, #0a0718 78%, #05030a 100%);
        box-shadow: inset 0 0 60px rgba(0,0,0,0.9);
    }}
    .aurora-band {{ position: absolute; top: 0; left: 0; width: 100%; height: 65%; filter: blur(26px); opacity: 0.55; mix-blend-mode: screen; pointer-events: none; }}
    .aurora-band.a1 {{ background: linear-gradient(180deg, rgba(110,220,170,0.60), transparent 75%); animation: auroraShift 8s ease-in-out infinite; }}
    .aurora-band.a2 {{ background: linear-gradient(180deg, rgba(167,139,250,0.55), transparent 70%); animation: auroraShift 11s ease-in-out infinite reverse; }}
    .aurora-band.a3 {{ background: linear-gradient(180deg, rgba(96,165,250,0.45), transparent 65%); animation: auroraShift 14s ease-in-out infinite; }}
    @keyframes auroraShift {{ 0%, 100% {{ transform: translateX(-6%) skewX(-4deg); opacity: 0.40; }} 50% {{ transform: translateX(6%) skewX(4deg); opacity: 0.75; }} }}
    .mountain-back, .mountain-front {{ position: absolute; bottom: 0; left: 0; width: 100%; pointer-events: none; }}
    .mountain-back {{ height: 48%; background: linear-gradient(180deg, #1a1230 0%, #0d081c 100%); clip-path: polygon(0% 100%, 0% 62%, 12% 42%, 22% 56%, 34% 32%, 46% 52%, 58% 28%, 70% 48%, 82% 34%, 92% 52%, 100% 42%, 100% 100%); }}
    .mountain-front {{ height: 32%; background: #04020a; clip-path: polygon(0% 100%, 0% 72%, 18% 52%, 32% 68%, 46% 44%, 60% 62%, 74% 48%, 88% 66%, 100% 56%, 100% 100%); }}
    .shooting-star {{ position: absolute; width: 110px; height: 2px; border-radius: 2px; background: linear-gradient(90deg, rgba(255,255,255,0) 0%, rgba(200,176,232,0.55) 55%, #ffffff 100%); transform: rotate(35deg); transform-origin: right center; opacity: 0; pointer-events: none; z-index: 4; }}
    .shooting-star::after {{ content: ""; position: absolute; right: -3px; top: -2px; width: 5px; height: 5px; background: #ffffff; border-radius: 50%; box-shadow: 0 0 10px #ffffff, 0 0 20px rgba(167,139,250,0.9); }}
    .shooting-star.ss1 {{ top: 10%; left: -15%; animation: shootStar 7s linear infinite; animation-delay: 0s; }}
    .shooting-star.ss2 {{ top: 22%; left: -15%; animation: shootStar 8s linear infinite; animation-delay: 2.4s; }}
    .shooting-star.ss3 {{ top: 4%; left: -15%; animation: shootStar 9s linear infinite; animation-delay: 4.8s; }}
    @keyframes shootStar {{ 0% {{ top: 8%; left: -15%; opacity: 0; }} 6% {{ opacity: 1; }} 45% {{ top: 58%; left: 85%; opacity: 1; }} 55% {{ top: 62%; left: 92%; opacity: 0; }} 100% {{ top: 62%; left: 92%; opacity: 0; }} }}
    .star {{ position: absolute; border-radius: 50%; background: #c8b0e8; box-shadow: 0 0 4px #c8b0e8; animation: twinkle 1.2s ease-in-out infinite; z-index: 2; }}
    .star.amber {{ background: #f0b860; box-shadow: 0 0 6px #f0b860; }}
    .star.purple {{ background: #a78bfa; box-shadow: 0 0 6px #a78bfa; }}
    .star.s1 {{ top: 12%; left: 8%; width: 2px; height: 2px; animation-delay: 0.0s; }}
    .star.s2 {{ top: 24%; left: 18%; width: 3px; height: 3px; animation-delay: 0.4s; }}
    .star.s3 {{ top: 8%; left: 32%; width: 2px; height: 2px; animation-delay: 0.8s; }}
    .star.s4 {{ top: 30%; left: 44%; width: 3px; height: 3px; animation-delay: 1.2s; }}
    .star.s5 {{ top: 14%; left: 56%; width: 2px; height: 2px; animation-delay: 0.3s; }}
    .star.s6 {{ top: 26%; left: 66%; width: 3px; height: 3px; animation-delay: 0.9s; }}
    .star.s7 {{ top: 6%; left: 78%; width: 2px; height: 2px; animation-delay: 1.5s; }}
    .star.s8 {{ top: 18%; left: 90%; width: 3px; height: 3px; animation-delay: 0.6s; }}
    .star.s9 {{ top: 36%; left: 24%; width: 2px; height: 2px; animation-delay: 1.8s; }}
    .star.s10 {{ top: 40%; left: 72%; width: 3px; height: 3px; animation-delay: 0.2s; }}
    .star.s11 {{ top: 4%; left: 50%; width: 2px; height: 2px; animation-delay: 1.0s; }}
    .star.s12 {{ top: 33%; left: 88%; width: 2px; height: 2px; animation-delay: 1.4s; }}
    @keyframes twinkle {{ 0%, 100% {{ opacity: 0.20; transform: scale(0.5); }} 50% {{ opacity: 1.0; transform: scale(1.5); }} }}

    @media (prefers-reduced-motion: reduce) {{
        .aurora-band, .shooting-star, .star, .preview-hero-badge {{ animation: none !important; }}
    }}

    .preview-hero {{ display: flex; align-items: center; gap: 12px; padding: 8px 14px; margin-bottom: 6px; background: linear-gradient(135deg, rgba(167,139,250,0.15) 0%, rgba(96,165,250,0.15) 100%); border-radius: 8px; border: 1px solid {border}; }}
    .preview-hero-title {{ font-size: 1.05rem; font-weight: 800; background: linear-gradient(90deg, #a78bfa, #60a5fa); -webkit-background-clip: text; -webkit-text-fill-color: transparent; background-clip: text; margin: 0; padding: 0; }}
    .preview-hero-badge {{ margin-left: auto; font-size: 0.62rem; font-weight: 800; padding: 3px 9px; background: transparent; border: 1px solid {accent}; color: {accent} !important; border-radius: 999px; letter-spacing: 0.08em; animation: pulseText 2.4s ease-in-out infinite; }}
    @keyframes pulseText {{ 0%, 100% {{ opacity: 1; }} 50% {{ opacity: 0.6; }} }}

    .nudge-header {{ font-size: 0.72rem !important; font-weight: 700; color: {text_soft} !important; text-transform: uppercase; letter-spacing: 0.06em; margin: 4px 0 2px 0 !important; display: block; }}

    .info-bar {{ display: block; width: 100%; padding: 4px 12px; margin: 2px 0 4px 0; background: {panel}; border-left: 3px solid {accent}; border-radius: 4px; font-size: 0.74rem; color: {text_soft} !important; letter-spacing: 0.02em; }}
    .warn-bar {{ display: block; width: 100%; padding: 4px 12px; margin: 2px 0 4px 0; background: {panel}; border-left: 3px solid #d97757; border-radius: 4px; font-size: 0.74rem; color: {text} !important; letter-spacing: 0.02em; }}

    div[data-testid="stCustomComponentV1"],
    iframe[title="cv_core.drag_crop.drag_crop_box"],
    iframe[title*="drag_crop"] {{ margin-top: 14px !important; margin-bottom: 6px !important; border-radius: 8px; overflow: hidden; }}
    .drag-hint, .drag-hint-text, div[class*="hint"] {{ margin-top: 12px !important; font-size: 0.78rem !important; color: {text_soft} !important; text-align: center !important; display: block !important; visibility: visible !important; opacity: 1 !important; }}

    div[data-testid="stAlert"] {{ padding: 4px 10px !important; margin: 0.15rem 0 !important; }}
    div[data-testid="stAlert"] p {{ font-size: 0.78rem !important; margin: 0 !important; }}

    {dark_extras}
    {colorful_extras}
    </style>
    """

st.markdown(_theme_css(st.session_state.theme_mode), unsafe_allow_html=True)

# ----------------------------------------------------------------------
# Sidebar
# ----------------------------------------------------------------------
with st.sidebar:
    st.markdown("### 🎨 Theme")
    theme_choice = st.radio(
        "Theme", ["Light","Gray","Dark","Colorful"],
        index=["Light","Gray","Dark","Colorful"].index(st.session_state.theme_mode)
        if st.session_state.theme_mode in ("Light","Gray","Dark","Colorful") else 1,
        format_func=lambda m: {"Light":"☀️ Light","Gray":"🌫️ Gray","Dark":"🌙 Dark","Colorful":"🎨 Colorful"}[m],
        horizontal=True, key="theme_radio_selector", label_visibility="collapsed")
    if theme_choice != st.session_state.theme_mode:
        st.session_state.theme_mode = theme_choice; st.rerun()

    st.markdown("---")
    st.markdown("**1. Upload**")
    uploaded = st.file_uploader("Composite images (passport + 3×4 + full body)",
                                 type=["jpg","jpeg","png"], accept_multiple_files=True, key=UPLOADER_KEY)
    if uploaded:
        new_names = [f.name for f in uploaded]
        old_cache = st.session_state.get(CACHE_KEY, [])
        old_names = [c["name"] for c in old_cache]
        if new_names != old_names:
            per_img_now = st.session_state.get("per_image", {})
            for old_n in old_names:
                if old_n not in new_names: per_img_now.pop(old_n, None)
            st.session_state[CACHE_KEY] = [{"name": f.name, "bytes": f.getvalue()} for f in uploaded]
            st.session_state[HASH_KEY] = {f.name: _file_hash(f.getvalue()) for f in uploaded}
        files = list(uploaded)
    elif st.session_state.get(CACHE_KEY):
        files = [_CachedUpload(c["name"], c["bytes"]) for c in st.session_state[CACHE_KEY]]
    else: files = []

    if files:
        hashes = st.session_state.get(HASH_KEY, {})
        if not hashes:
            hashes = {f.name: _file_hash(f.getvalue()) for f in files}
            st.session_state[HASH_KEY] = hashes
        seen = {}; dup_pairs = []
        for n in [f.name for f in files]:
            h = hashes.get(n)
            if h is None: continue
            if h in seen: dup_pairs.append((seen[h], n))
            else: seen[h] = n
        if dup_pairs:
            for orig, dup in dup_pairs: st.warning(f"⚠️ **{dup}** looks identical to **{orig}**")
        if st.button("🗑️ Clear all uploaded", use_container_width=True):
            for k in [UPLOADER_KEY, CACHE_KEY, HASH_KEY, "per_image", "active_file",
                      "boxes", "extracted", "cv_data", "approved",
                      "_auto_read_done_for", "_pending_switch_to", PREVIEW_KEY]:
                st.session_state.pop(k, None)
            st.rerun()

    # ── Session save / load ──
    with st.expander("💾 Save / Load session"):
        st.caption("Save the full working state (all crops, approved CVs, settings) to a JSON file — reload later to continue.")
        if st.session_state.get("per_image"):
            st.download_button(
                "📥 Download session (JSON)",
                data=_session_to_json(),
                file_name=f"cv_session_{datetime.today():%Y-%m-%d_%H%M}.json",
                mime="application/json",
                use_container_width=True,
                key="dl_session_json",
            )
        else:
            st.caption("_Nothing to save yet — upload some images first._")
        loaded_file = st.file_uploader("Load session JSON", type=["json"], key="load_session_uploader")
        if loaded_file is not None and st.button("🔄 Restore from file", use_container_width=True, key="restore_session_btn"):
            try:
                _load_session_json(loaded_file.read().decode("utf-8"))
                st.success("Session restored.")
                st.rerun()
            except Exception as _e:
                st.error(f"Could not restore: {_e}")

    st.markdown("---")
    st.markdown("**2. Image enlargement (k)**")
    st.caption("Base sizes (inches):\n• 3×4 → 1.10 × 1.30\n• Full → 1.85 × 3.30\n• Passport → 1.38 × 1.59\nEach is multiplied by its own k below.")

    st.markdown('<span class="k-row-label">3×4 photo — k</span>', unsafe_allow_html=True)
    _cur = float(st.session_state.get("k_3x4", DEFAULT_K_3X4))
    _idx = K_OPTIONS_SMALL.index(_cur) if _cur in K_OPTIONS_SMALL else K_OPTIONS_SMALL.index(1.0)
    st.markdown('<div class="k-row">', unsafe_allow_html=True)
    k_3x4_choice = st.radio("3x4 k", K_OPTIONS_SMALL, index=_idx, horizontal=True,
                             format_func=lambda v: f"{v}", key="k_3x4_radio", label_visibility="collapsed")
    st.markdown('</div>', unsafe_allow_html=True)
    st.session_state.k_3x4 = float(k_3x4_choice)

    st.markdown('<span class="k-row-label">Full-body photo — k</span>', unsafe_allow_html=True)
    _cur = float(st.session_state.get("k_full", DEFAULT_K_FULL))
    _idx = K_OPTIONS_FULL.index(_cur) if _cur in K_OPTIONS_FULL else K_OPTIONS_FULL.index(1.0)
    st.markdown('<div class="k-row">', unsafe_allow_html=True)
    k_full_choice = st.radio("full k", K_OPTIONS_FULL, index=_idx, horizontal=True,
                              format_func=lambda v: f"{v}", key="k_full_radio", label_visibility="collapsed")
    st.markdown('</div>', unsafe_allow_html=True)
    st.session_state.k_full = float(k_full_choice)

    st.markdown('<span class="k-row-label">Passport photo — k</span>', unsafe_allow_html=True)
    _cur = float(st.session_state.get("k_passport", DEFAULT_K_PASSPORT))
    _idx = K_OPTIONS.index(_cur) if _cur in K_OPTIONS else K_OPTIONS.index(3.0)
    st.markdown('<div class="k-row">', unsafe_allow_html=True)
    k_passport_choice = st.radio("passport k", K_OPTIONS, index=_idx, horizontal=True,
                                  format_func=lambda v: f"{v}", key="k_passport_radio", label_visibility="collapsed")
    st.markdown('</div>', unsafe_allow_html=True)
    st.session_state.k_passport = float(k_passport_choice)

    st.markdown("---")
    st.markdown("**3. Shortcuts**")
    st.session_state.enable_shortcuts = st.checkbox(
        "⌨️ Enable keyboard shortcuts",
        value=st.session_state.enable_shortcuts,
        help="Enter = Approve · P = Preview · N = Next image (works when focus isn't in a text field)."
    )

    st.markdown("---")
    st.markdown("**4. OCR settings**")
    fast_mode = st.checkbox("⚡ Fast OCR mode", value=True)
    use_easy = st.checkbox("Use EasyOCR (slower)", value=False)
    use_gemini = st.checkbox("Use Gemini for missing fields", value=False)
    gemini_keys = st.text_input("Gemini API key(s) — one per line", value=os.getenv("GEMINI_API_KEYS", ""), type="password")
    gemini_model = st.text_input("Gemini model", value="gemini-3.7-flash")
    amharic = st.checkbox("Amharic assist (translate to English)")

    st.markdown("---")
    if st.button("🔄 System check", use_container_width=True):
        e = get_engine_status()
        st.write("Tesseract:", "✅" if e.tesseract_executable else "❌")
        st.write("EasyOCR:", "✅" if e.easyocr else "❌")
        st.write("MRZ package:", "✅" if e.mrz_package else "❌")
        st.write("PassportEye:", "✅" if e.passporteye else "❌")
        if e.languages: st.caption("Languages: " + ", ".join(e.languages))

# ----------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------
st.markdown("<h1 class='hero-title'>📋 CV Automation</h1>", unsafe_allow_html=True)
st.markdown("<p class='hero-sub'>Upload images, configure each one, approve them one by one, then generate all CVs.</p>", unsafe_allow_html=True)

if not files:
    st.markdown('<div style="height: 6vh;"></div>', unsafe_allow_html=True)
    st.info("👈 Upload one or more images using the sidebar to get started.")
    st.stop()

file_names = [f.name for f in files]
per_img = _per_image_dict()
for n in file_names: per_img.setdefault(n, _blank_image_slot())

pending_switch = st.session_state.pop("_pending_switch_to", None)
if pending_switch and pending_switch in file_names: _switch_to(pending_switch)
if "active_file" not in st.session_state or st.session_state.active_file not in file_names:
    unapproved = [n for n in file_names if not per_img[n]["approved"]]
    st.session_state.active_file = unapproved[0] if unapproved else file_names[0]

active_name = st.session_state.active_file
slot = per_img[active_name]
current_file = next(f for f in files if f.name == active_name)

with st.sidebar:
    st.markdown("---")
    approved_count = sum(1 for n in file_names if per_img[n]["approved"])
    st.markdown(f"**5. Images ({approved_count}/{len(file_names)} approved)**")
    for n in file_names:
        s = per_img[n]
        cd = s.get("cv_data") or {}
        ex = s.get("extracted") or {}
        display_name = cd.get("NAME") or ex.get("NAME") or ""
        issue = _name_issue_reason(display_name)
        mark = "✅" if s["approved"] else "⏳"
        warn_mark = " ⚠️" if (issue and not s["approved"]) else ""
        is_active = (n == active_name)
        row1, row2 = st.columns([0.75, 0.25], gap="small")
        with row1:
            if st.button(f"{mark}{warn_mark} {n}", key=f"pick_{n}", use_container_width=True,
                         type="primary" if is_active else "secondary", disabled=is_active):
                st.session_state._pending_switch_to = n; st.rerun()
        with row2:
            if s["approved"]:
                if st.button("↩", key=f"unapp_{n}", help="Unapprove this image to edit it again",
                             use_container_width=True):
                    s["approved"] = False
                    if n == active_name:
                        st.session_state.approved = False
                    st.rerun()
        st.caption(f"   {s['agency']} · {s['validity_years']}y · {s['level']} · {s['religion']}")
        if issue and not s["approved"]: st.caption(f"   ⚠️ {issue}")

total = len(file_names)
approved_count = sum(1 for n in file_names if per_img[n]["approved"])
with st.container(key="sticky_progress"):
    st.progress(approved_count / total, text=f"✅ Approved {approved_count} of {total} images")

# ── Keyboard shortcuts (injected into parent doc) ──
if st.session_state.get("enable_shortcuts", True):
    components.html("""
    <script>
    (function() {
      const doc = window.parent.document;
      if (doc._cv_kb_attached) return;
      doc._cv_kb_attached = true;
      doc.addEventListener('keydown', function(e) {
        const t = e.target;
        if (t && (t.tagName === 'INPUT' || t.tagName === 'TEXTAREA' || t.isContentEditable)) return;
        const key = (e.key || '').toLowerCase();
        const btns = Array.from(doc.querySelectorAll('button'));
        if (key === 'enter') {
          for (const b of btns) {
            const txt = (b.innerText || '').trim();
            if (txt.startsWith('✅ Approve')) { b.click(); e.preventDefault(); return; }
          }
        } else if (key === 'p') {
          for (const b of btns) {
            const txt = (b.innerText || '').trim();
            if (txt.includes('Preview current')) { b.click(); e.preventDefault(); return; }
          }
        } else if (key === 'n') {
          const picks = btns.filter(b => {
            const txt = (b.innerText || '').trim();
            return txt.startsWith('⏳') || txt.startsWith('✅') || txt.startsWith('⚠️');
          });
          const cur = picks.findIndex(b => b.disabled);
          if (cur >= 0 && picks.length > 1) {
            for (let i = 1; i <= picks.length; i++) {
              const cand = picks[(cur + i) % picks.length];
              if (!cand.disabled) { cand.click(); e.preventDefault(); return; }
            }
          }
        }
      });
    })();
    </script>
    """, height=0)

st.markdown(f'<div class="config-title">📌 Configuring: <span style="color:#3b82f6">{active_name}</span></div>', unsafe_allow_html=True)

st.session_state.setdefault("cfg_agency", slot["agency"])
st.session_state.setdefault("cfg_level", slot["level"])
st.session_state.setdefault("cfg_validity", slot["validity_years"])
st.session_state.setdefault("cfg_religion", slot["religion"])

c1, c2, c3, c4 = st.columns([1.1, 1.1, 1.0, 1.0], gap="small")
with c1: cfg_agency = st.selectbox("Agency", REAL_AGENCIES + [BOTH_LABEL], key="cfg_agency")
with c2: cfg_level = st.selectbox("Experience", EXPERIENCE_LEVELS, key="cfg_level")
with c3: cfg_validity = st.selectbox("Passport term", VALIDITY_OPTIONS, format_func=lambda v: f"{v} years", key="cfg_validity")
with c4: cfg_religion = st.selectbox("CV religion", RELIGION_OPTIONS, key="cfg_religion")

slot["agency"] = cfg_agency; slot["level"] = cfg_level
slot["validity_years"] = int(cfg_validity); slot["religion"] = cfg_religion

# ── Apply to all button ──
_apply_col1, _apply_col2 = st.columns([1, 1], gap="small")
with _apply_col1:
    if st.button("📋 Apply these 4 settings to ALL images", key="apply_cfg_all",
                 use_container_width=True,
                 help="Set the same Agency / Experience / Passport term / Religion for every image in the batch."):
        for n in file_names:
            per_img[n]["agency"] = cfg_agency
            per_img[n]["level"] = cfg_level
            per_img[n]["validity_years"] = int(cfg_validity)
            per_img[n]["religion"] = cfg_religion
        st.success(f"Applied to all {len(file_names)} image(s).")
        st.rerun()

# ── Copy from previous ──
_prev_name = None
try:
    _idx_active = file_names.index(active_name)
    if _idx_active > 0: _prev_name = file_names[_idx_active - 1]
except ValueError: pass

with _apply_col2:
    if _prev_name and per_img[_prev_name].get("cv_data"):
        if st.button(f"📋 Copy fields from '{_prev_name}'", key=f"copy_prev_{active_name}",
                     use_container_width=True,
                     help="Copy Name / Passport / Dates from the previous image."):
            prev_slot = per_img[_prev_name]
            slot["cv_data"] = dict(prev_slot["cv_data"])
            if prev_slot.get("extracted"): slot["extracted"] = dict(prev_slot["extracted"])
            st.session_state["_force_date_reseed"] = True
            st.rerun()

prev_validity = st.session_state.get("_prev_cfg_validity")
if prev_validity is not None and prev_validity != cfg_validity:
    st.session_state["_force_date_reseed"] = True
st.session_state["_prev_cfg_validity"] = cfg_validity

prev_level = st.session_state.get("_prev_cfg_level")
if prev_level is not None and prev_level != cfg_level:
    for k in ("exp_position_in","exp_country_select","exp_country_manual","exp_years_exp"):
        st.session_state.pop(k, None)
st.session_state["_prev_cfg_level"] = cfg_level

cur_validity = int(cfg_validity); cur_level = cfg_level; cur_agency = cfg_agency; cur_religion = cfg_religion

# ── Load image with per-image rotation ──
raw_img = Image.open(BytesIO(current_file.getvalue())).convert("RGB")
_rot = int(slot.get("rotation", 0)) % 360
img = raw_img.rotate(-_rot, expand=True) if _rot else raw_img
w, h = img.size
if slot.get("boxes") is None: slot["boxes"] = auto_detect_boxes(img)
boxes = slot["boxes"]; st.session_state.boxes = boxes

if (not slot.get("extracted") and not slot.get("approved")
    and st.session_state.get("_auto_read_done_for") != active_name):
    st.session_state["_auto_read_done_for"] = active_name
    try:
        auto_passport_crop = img.crop(boxes["passport"])
        auto_ocr_crop = trim_whitespace(img.crop(boxes["ocr"])) if boxes.get("ocr") else None
        p_buf, o_buf = BytesIO(), BytesIO()
        auto_passport_crop.save(p_buf, format="JPEG", quality=92)
        if auto_ocr_crop is not None: auto_ocr_crop.save(o_buf, format="JPEG", quality=92)
        with st.spinner(f"⚡ Reading passport for {active_name}…"):
            extracted = cached_smart_read(p_buf.getvalue(), o_buf.getvalue() if auto_ocr_crop is not None else b"",
                                           use_easy, use_gemini, gemini_keys, gemini_model,
                                           amharic, fast_mode, validity_years=cur_validity)
            if not extracted.get("NAME"):
                fn = _name_from_filename(active_name)
                if fn: extracted["NAME"] = fn; extracted["_name_from_filename"] = True
            slot["extracted"] = extracted; st.session_state.extracted = extracted
    except Exception as e:
        slot["extracted"] = {"NAME": _name_from_filename(active_name), "PASSPORT_NO": "",
                              "DOB": "", "ISSUE_DATE": "", "EXPIRY_DATE": "", "HOME_ADDRESS": "", "AGE": "",
                              "_warning": f"Automatic OCR failed: {e}", "_mrz": {}, "_ocr_text": "", "_engine": "error"}
        st.session_state.extracted = slot["extracted"]

# ======================================================================
# STEP 1 — Crop regions
# ======================================================================
st.markdown('<span class="step-tag step-1">STEP 1</span>', unsafe_allow_html=True)

active_key = st.session_state.get("active_crop_region", CROP_KEYS[0])
st.markdown('<div class="live-crop-head"><div class="live-crop-label">Live full-image crop view</div></div>', unsafe_allow_html=True)

crop_controls, live_preview = st.columns([0.85, 1.6], gap="medium", vertical_alignment="top")
new_box = boxes[active_key]; undo_key = f"_prev_box_{active_name}"

with crop_controls:
    col_buttons = st.columns(2)
    if col_buttons[0].button("🔄 Auto-detect all", use_container_width=True):
        st.session_state[undo_key] = tuple(boxes[active_key])
        slot["boxes"] = auto_detect_boxes(img); st.session_state.boxes = slot["boxes"]
        st.session_state.canvas_version = st.session_state.get("canvas_version", 0) + 1; st.rerun()
    if col_buttons[1].button("↩️ Reset all", use_container_width=True):
        st.session_state[undo_key] = tuple(boxes[active_key])
        slot["boxes"] = default_boxes(w, h); st.session_state.boxes = slot["boxes"]
        st.session_state.canvas_version = st.session_state.get("canvas_version", 0) + 1; st.rerun()

    st.markdown('<span class="k-row-label">Rotate source image</span>', unsafe_allow_html=True)
    rc = st.columns(3)
    if rc[0].button("↺ 90°", use_container_width=True, help="Rotate 90° counter-clockwise"):
        slot["rotation"] = (int(slot.get("rotation", 0)) + 90) % 360
        slot["boxes"] = None; st.rerun()
    if rc[1].button("↻ 90°", use_container_width=True, help="Rotate 90° clockwise"):
        slot["rotation"] = (int(slot.get("rotation", 0)) - 90) % 360
        slot["boxes"] = None; st.rerun()
    if rc[2].button("⟲ Reset", use_container_width=True, help="Reset rotation to 0°"):
        slot["rotation"] = 0; slot["boxes"] = None; st.rerun()
    if slot.get("rotation", 0):
        st.caption(f"Current rotation: {slot['rotation']}°")

    st.caption(CROP_HINTS[active_key])
    new_box = crop_with_numbers(img, active_key, boxes[active_key], w, h)

    col_action = st.columns(3)
    if col_action[0].button("Auto", use_container_width=True):
        st.session_state[undo_key] = tuple(boxes[active_key])
        boxes[active_key] = auto_detect_boxes(img)[active_key]
        slot["boxes"] = boxes; st.session_state.boxes = boxes
        st.session_state.canvas_version = st.session_state.get("canvas_version", 0) + 1; st.rerun()
    if col_action[1].button("Reset", use_container_width=True):
        st.session_state[undo_key] = tuple(boxes[active_key])
        boxes[active_key] = default_boxes(w, h)[active_key]
        slot["boxes"] = boxes; st.session_state.boxes = boxes
        st.session_state.canvas_version = st.session_state.get("canvas_version", 0) + 1; st.rerun()
    if col_action[2].button("↩ Undo", use_container_width=True):
        prev = st.session_state.get(undo_key)
        if prev:
            boxes[active_key] = tuple(prev); slot["boxes"] = boxes; st.session_state.boxes = boxes
            st.session_state.canvas_version = st.session_state.get("canvas_version", 0) + 1; st.rerun()

    st.markdown("""
    <div class="crop-help-box">
      <div class="crop-help-title">🎬 How to crop — 3 steps</div>
      <div class="crop-help-body">
        <b>1.</b> Pick a region with the round buttons on the right.<br>
        <b>2.</b> Drag the colored box on the photo to set the crop.<br>
        <b>3.</b> Use <b>Auto</b> to snap, <b>Reset</b> for default, <b>↩ Undo</b> to revert.
      </div>
    </div>
    """, unsafe_allow_html=True)

with live_preview:
    st.subheader("✂️ Crop regions")
    st.markdown(f'<div class="crop-caption">Drag on the photo to set <b>{CROP_LABELS[active_key]}</b>.</div>', unsafe_allow_html=True)
    st.markdown('<div class="crop-legend">Blue: Passport · Orange: 3×4 · Green: Full body · Purple: OCR</div>', unsafe_allow_html=True)
    st.markdown('<div class="crop-region-row">', unsafe_allow_html=True)
    active_key = st.radio("Crop region", CROP_KEYS, format_func=lambda key: CROP_LABELS[key],
                          key="active_crop_region", label_visibility="collapsed", horizontal=True)
    st.markdown('</div>', unsafe_allow_html=True)
    new_box = crop_with_canvas(img, active_key, boxes[active_key], canvas_width=CROP_PREVIEW_WIDTH)

if any(abs(new_box[i] - boxes[active_key][i]) > 2 for i in range(4)):
    st.session_state[undo_key] = tuple(boxes[active_key])
    boxes[active_key] = new_box; slot["boxes"] = boxes; st.session_state.boxes = boxes; st.rerun()

slot["boxes"] = boxes; st.session_state.boxes = boxes

# ======================================================================
# STEP 2 — Passport extraction
# ======================================================================
st.markdown('<span class="step-tag step-2">STEP 2</span>', unsafe_allow_html=True)

_hdr_l, _hdr_r = st.columns([0.55, 3.45], vertical_alignment="center")
with _hdr_l:
    smart_read_clicked = st.button("🚀 Smart read", type="primary", use_container_width=True,
                                    help="Reads the OCR crop and fills the fields on the right.")
with _hdr_r:
    st.markdown('<span class="ocr-header">📖 Passport extraction</span>', unsafe_allow_html=True)

ex = st.session_state.get("extracted") or {"NAME":"","PASSPORT_NO":"","DOB":"","ISSUE_DATE":"","EXPIRY_DATE":"","HOME_ADDRESS":"","AGE":"","_warning":""}
warn = ex.get("_warning", "")
if warn:
    warn = warn.replace("ISSUE_DATE, ", "").replace(", ISSUE_DATE", "").replace("ISSUE_DATE", "")
    if warn.strip() == "Please fill in manually: .": warn = ""
if ex.get("_name_from_filename"):
    warn = (warn + " " if warn else "") + "Name guessed from filename — please verify."

if warn.strip():
    st.markdown(f'<div class="info-bar">ℹ️ {warn.strip()}</div>', unsafe_allow_html=True)

passport_crop = img.crop(boxes["passport"])
ocr_crop = trim_whitespace(img.crop(boxes["ocr"])) if boxes.get("ocr") else None
ocr_preview_col, fields_col = st.columns([1.0, 1.35], gap="medium", vertical_alignment="top")

with fields_col:
    force_reseed = st.session_state.pop("_force_date_reseed", False)
    if force_reseed:
        for _mk in ("dob_mode_radio","issue_mode_radio","expiry_mode_radio"): st.session_state.pop(_mk, None)
        for _mk in ("_prev_dob_mode","_prev_issue_mode","_prev_expiry_mode","_last_dob_for_suggestion"): st.session_state.pop(_mk, None)
        saved_cv = slot.get("cv_data") or {}
        st.session_state["full_name_in"] = saved_cv.get("NAME") or ex.get("NAME", "")
        st.session_state["passport_no_in"] = saved_cv.get("PASSPORT_NO") or ex.get("PASSPORT_NO", "")
        st.session_state["place_of_birth_in"] = saved_cv.get("HOME_ADDRESS") or ex.get("HOME_ADDRESS", "")
        st.session_state["marital_select"] = saved_cv.get("MARITAL_STATUS", "SINGLE")
        st.session_state["children_select"] = saved_cv.get("NO_OF_CHILDREN", "NIL")
        st.session_state.setdefault("relative_name", saved_cv.get("RELATIVE", ""))

    st.markdown('<div class="compact-card">', unsafe_allow_html=True)
    st.markdown('<div class="details-header">👤 Personal Information</div>', unsafe_allow_html=True)
    c1, c2 = st.columns(2, gap="large")
    with c1: name = st.text_input("Full name", key="full_name_in")
    with c2: passport = st.text_input("Passport No.", key="passport_no_in")
    st.markdown('<div class="field-row-gap"></div>', unsafe_allow_html=True)
    c1, c2 = st.columns(2, gap="large")
    with c1: pob = st.text_input("Place of birth", key="place_of_birth_in")
    with c2: relative = st.text_input("Relative", key="relative_name")

    st.markdown('<hr class="section-divider">', unsafe_allow_html=True)
    st.markdown('<div class="details-header">📅 Passport Details</div>', unsafe_allow_html=True)

    saved_cv = slot.get("cv_data") or {}
    seed_dob = saved_cv.get("DOB") or ex.get("DOB", "")
    seed_issue = saved_cv.get("ISSUE_DATE") or ex.get("ISSUE_DATE", "")
    seed_expiry = saved_cv.get("EXPIRY_DATE") or ex.get("EXPIRY_DATE", "")

    prev_dob_mode = st.session_state.get("_prev_dob_mode")
    hdr = st.columns([1.7, 1.3], gap="small", vertical_alignment="center")
    with hdr[0]: st.markdown('<div class="date-section-title">🎂 Date of birth</div>', unsafe_allow_html=True)
    with hdr[1]:
        st.markdown('<div class="mode-pill">', unsafe_allow_html=True)
        dob_mode = st.radio("dob_mode", ["Auto","Manual"], horizontal=True, key="dob_mode_radio", label_visibility="collapsed")
        st.markdown('</div>', unsafe_allow_html=True)
    if force_reseed or prev_dob_mode is None or prev_dob_mode != dob_mode:
        _seed_date_widgets("Date of birth", seed_dob, 1981, 2007)
        st.session_state.pop("_last_dob_for_suggestion", None)
    st.session_state["_prev_dob_mode"] = dob_mode
    if dob_mode == "Auto":
        if seed_dob:
            dob = seed_dob
            st.markdown(f'<div class="derived-date"><span class="dd-label">Value</span><span class="dd-value">{seed_dob}</span><span class="dd-hint">from passport</span></div>', unsafe_allow_html=True)
        else:
            dob = ""; st.caption("⚠️ No date of birth detected — switch to Manual.")
    else: dob = date_picker("Date of birth", 1981, 2007)

    st.markdown('<div class="date-block-gap"></div>', unsafe_allow_html=True)

    prev_issue_mode = st.session_state.get("_prev_issue_mode")
    hdr = st.columns([1.7, 1.3], gap="small", vertical_alignment="center")
    with hdr[0]: st.markdown('<div class="date-section-title">📅 Issue date</div>', unsafe_allow_html=True)
    with hdr[1]:
        st.markdown('<div class="mode-pill">', unsafe_allow_html=True)
        issue_mode = st.radio("issue_mode", ["Auto","Manual"], horizontal=True, key="issue_mode_radio", label_visibility="collapsed")
        st.markdown('</div>', unsafe_allow_html=True)
    if force_reseed or prev_issue_mode is None or prev_issue_mode != issue_mode:
        _seed_date_widgets("Issue date", seed_issue, 2016, 2030)
    st.session_state["_prev_issue_mode"] = issue_mode
    if issue_mode == "Auto":
        derived_issue = _subtract_passport_years(seed_expiry, cur_validity) if seed_expiry else ""
        if derived_issue:
            issue = derived_issue
            st.markdown(f'<div class="derived-date"><span class="dd-label">Value</span><span class="dd-value">{derived_issue}</span><span class="dd-hint">expiry − {cur_validity}y + 1d</span></div>', unsafe_allow_html=True)
        else:
            issue = ""; st.caption("⚠️ No expiry date yet — switch to Manual or set expiry below.")
    else: issue = date_picker("Issue date", 2016, 2030)

    st.markdown('<div class="date-block-gap"></div>', unsafe_allow_html=True)

    prev_expiry_mode = st.session_state.get("_prev_expiry_mode")
    hdr = st.columns([1.7, 1.3], gap="small", vertical_alignment="center")
    with hdr[0]: st.markdown('<div class="date-section-title">📅 Expiry date</div>', unsafe_allow_html=True)
    with hdr[1]:
        st.markdown('<div class="mode-pill">', unsafe_allow_html=True)
        expiry_mode = st.radio("expiry_mode", ["Auto","Manual"], horizontal=True, key="expiry_mode_radio", label_visibility="collapsed")
        st.markdown('</div>', unsafe_allow_html=True)
    if force_reseed or prev_expiry_mode is None or prev_expiry_mode != expiry_mode:
        _seed_date_widgets("Expiry date", seed_expiry, 2021, 2038)
    st.session_state["_prev_expiry_mode"] = expiry_mode
    if expiry_mode == "Auto":
        derived_expiry = _add_passport_years(issue, cur_validity) if issue else ""
        if derived_expiry:
            expiry = derived_expiry
            st.markdown(f'<div class="derived-date"><span class="dd-label">Value</span><span class="dd-value">{derived_expiry}</span><span class="dd-hint">issue + {cur_validity}y − 1d</span></div>', unsafe_allow_html=True)
        else:
            expiry = ""; st.caption("⚠️ No issue date yet — switch to Manual or set issue above.")
    else: expiry = date_picker("Expiry date", 2021, 2038)

    # ── Date sanity warnings ──
    _dw = _date_sanity_warnings(dob, issue, expiry)
    for _w in _dw:
        st.markdown(f'<div class="warn-bar">⚠️ {_w}</div>', unsafe_allow_html=True)

    if dob and dob != st.session_state.get("_last_dob_for_suggestion"):
        st.session_state["_last_dob_for_suggestion"] = dob
        sug_m, sug_c = suggest_marital_children(_age_from_dob(dob))
        st.session_state["marital_select"] = sug_m; st.session_state["children_select"] = sug_c

    st.markdown('<hr class="section-divider">', unsafe_allow_html=True)
    notes_key = f"notes_visible_{active_name}"
    st.session_state.setdefault(notes_key, False)
    if st.button(("📝 Hide note" if st.session_state[notes_key] else "📝 Add note"),
                 key=f"notes_btn_{active_name}", use_container_width=False):
        st.session_state[notes_key] = not st.session_state[notes_key]; st.rerun()
    if st.session_state[notes_key]:
        notes_val = st.text_area("📝 Notes (optional)", value=slot.get("notes", ""),
                                  key=f"img_notes_{active_name}",
                                  placeholder="e.g. Called customer — address changed",
                                  height=68)
        slot["notes"] = notes_val

    st.markdown('</div>', unsafe_allow_html=True)

with ocr_preview_col:
    st.markdown("🔍 OCR crop")
    if ocr_crop is not None: st.image(ocr_crop, caption=f"{ocr_crop.width} × {ocr_crop.height} px", width=440)
    else: st.warning("No OCR crop selected.")

    # ── Raw OCR text (collapsible) ──
    with st.expander("🔤 Raw OCR text", expanded=False):
        raw_text = (ex or {}).get("_ocr_text", "") or ""
        engine = (ex or {}).get("_engine", "")
        if engine: st.caption(f"Engine: {engine}")
        if raw_text.strip():
            st.code(raw_text, language=None)
        else:
            st.caption("_No OCR text captured yet. Click “Smart read” to run OCR._")

    if smart_read_clicked:
        with st.spinner("Reading passport..."):
            p_buf, o_buf = BytesIO(), BytesIO()
            passport_crop.save(p_buf, format="JPEG")
            if ocr_crop is not None: ocr_crop.save(o_buf, format="JPEG")
            extracted = cached_smart_read(p_buf.getvalue(), o_buf.getvalue() if ocr_crop is not None else b"",
                                           use_easy, use_gemini, gemini_keys, gemini_model,
                                           amharic, fast_mode, validity_years=cur_validity)
            if not extracted.get("NAME"):
                fn = _name_from_filename(active_name)
                if fn: extracted["NAME"] = fn; extracted["_name_from_filename"] = True
            slot["extracted"] = extracted; st.session_state.extracted = extracted
        st.session_state["_force_date_reseed"] = True; st.rerun()

    st.markdown('<hr class="section-divider">', unsafe_allow_html=True)
    st.markdown('<div class="details-header">👤 Additional Details</div>', unsafe_allow_html=True)
    marital = st.selectbox("Marital", ["SINGLE","MARRIED"], key="marital_select")
    children = st.selectbox("Children", ["NIL","1","2","3","4"], key="children_select")
    religion = st.selectbox("Religion", RELIGION_OPTIONS, index=0 if cur_religion == "MUSLIM" else 1)

    if cur_level == "Experienced":
        st.markdown('<hr class="section-divider">', unsafe_allow_html=True)
        st.markdown('<div class="details-header">💼 Experience Details</div>', unsafe_allow_html=True)
        st.session_state.setdefault("exp_position_in", slot.get("position") or DEFAULT_POSITION)
        position_value = st.text_input("Position", key="exp_position_in", placeholder=DEFAULT_POSITION).strip().upper()
        if not position_value: position_value = DEFAULT_POSITION
        _seed_country(slot.get("country"), "exp_country_select", "exp_country_manual")
        country_choice = st.selectbox("Country", COUNTRY_PRESETS + [OTHER_LABEL], key="exp_country_select")
        if country_choice == OTHER_LABEL: country_value = st.text_input("Country (type manually)", key="exp_country_manual").strip()
        else: country_value = country_choice
        st.session_state.setdefault("exp_years_exp", int(slot.get("years_exp", 2)))
        years_exp = st.selectbox("Years experience", YEARS_RANGE, key="exp_years_exp")
        slot["position"] = position_value; slot["country"] = country_value; slot["years_exp"] = int(years_exp)
    else:
        position_value = DEFAULT_POSITION; country_value = "NIL"; years_exp = None

st.markdown('<div class="details-space"></div>', unsafe_allow_html=True)

with st.container(key="approve_bar"):
    btn_a, btn_b = st.columns(2, gap="medium")
    with btn_a: preview_clicked = st.button("👁 Preview current CV", use_container_width=True)
    with btn_b:
        approve_clicked = st.button("✅ Approve & Next" if approved_count < total - 1 else "✅ Approve (last one)",
                                     type="primary", use_container_width=True)

def _build_preview_cv():
    lvl = cur_level
    yrs = f"{years_exp} YEARS" if lvl == "Experienced" else "NIL"
    return {"NAME": _clean_passport_name(name.strip().upper()), "PASSPORT_NO": passport.strip().upper(),
            "ISSUE_DATE": issue, "EXPIRY_DATE": expiry, "DOB": dob, "HOME_ADDRESS": pob.strip().upper(),
            "RELATIVE": relative.strip().upper(), "RELIGION": religion, "MARITAL_STATUS": marital,
            "NO_OF_CHILDREN": children, "AGE": _age_from_dob(dob), "POSITION": position_value,
            "COUNTRY": country_value if lvl == "Experienced" else "NIL", "YEARS_EXP": yrs,
            "YEARS_EXP2": yrs if lvl == "Experienced" else "FIRST TIME"}

if preview_clicked:
    if not name.strip(): st.error("Please fill in the Full name at least.")
    else:
        preview_cv = _build_preview_cv()
        pdf_bytes = _preview_pdf_bytes(preview_cv, cur_agency, cur_level, active_name, files, boxes,
                                        rotation=int(slot.get("rotation", 0)))
        if pdf_bytes:
            st.session_state[PREVIEW_KEY] = {"file": active_name, "cv": preview_cv, "pdf_bytes": pdf_bytes}
        else: st.error("Preview could not be generated.")

# ----------------------------------------------------------------------
# Preview & Edit side-by-side
# ----------------------------------------------------------------------
pv = st.session_state.get(PREVIEW_KEY)
if pv and pv.get("file") == active_name and pv.get("pdf_bytes"):
    st.markdown('<hr class="section-divider">', unsafe_allow_html=True)
    st.markdown("""
    <div class="preview-hero">
      <span class="preview-hero-title">🖼️ CV Preview &amp; Edit</span>
      <span class="preview-hero-badge">LIVE · EDITABLE</span>
    </div>
    """, unsafe_allow_html=True)

    pdf_col, edit_col = st.columns([1.15, 1.0], gap="medium", vertical_alignment="top")
    with pdf_col:
        st.markdown("**Live preview**")
        try:
            import fitz
            with fitz.open(stream=pv["pdf_bytes"], filetype="pdf") as doc:
                for idx in range(doc.page_count):
                    page = doc.load_page(idx)
                    pix = page.get_pixmap(matrix=fitz.Matrix(1.3, 1.3), alpha=False)
                    st.image(pix.tobytes("png"), caption=f"Page {idx+1}", use_container_width=True)
        except ImportError: st.warning("PyMuPDF (fitz) not installed — cannot preview PDF.")

    with edit_col:
        with st.expander("🎯 Adjust images with mouse (drag sliders)", expanded=False):
            st.caption("Drag the sliders to shift each photo left/right or up/down inside its crop box.")
            nudge_cols = st.columns(3, gap="small")
            new_boxes_edit = dict(boxes)
            for i, ck in enumerate(["passport", "3x4", "full"]):
                with nudge_cols[i]:
                    st.markdown(f'<span class="nudge-header">{CROP_LABELS[ck]}</span>', unsafe_allow_html=True)
                    l0, t0, r0, b0 = boxes[ck]
                    dx = st.slider("↔", -100, 100, 0, step=2,
                                    key=f"pv_nudge_dx_{ck}_{active_name}",
                                    label_visibility="collapsed")
                    dy = st.slider("↕", -100, 100, 0, step=2,
                                    key=f"pv_nudge_dy_{ck}_{active_name}",
                                    label_visibility="collapsed")
                    new_boxes_edit[ck] = clamp_box(
                        (l0 + dx, t0 + dy, r0 + dx, b0 + dy), w, h
                    )
            new_boxes_edit["ocr"] = boxes.get("ocr")
            if st.button("✅ Apply image shifts", key=f"pv_nudge_apply_{active_name}",
                         use_container_width=True, type="primary"):
                slot["boxes"] = new_boxes_edit
                st.session_state.boxes = new_boxes_edit
                nb = _preview_pdf_bytes(pv["cv"], cur_agency, cur_level, active_name, files, new_boxes_edit,
                                         rotation=int(slot.get("rotation", 0)))
                if nb:
                    st.session_state[PREVIEW_KEY]["pdf_bytes"] = nb
                st.session_state.canvas_version = st.session_state.get("canvas_version", 0) + 1
                st.rerun()

        st.markdown("**✏️ Edit fields**")
        pc = dict(pv["cv"])

        with st.expander("👤 Name · Passport · DOB · Place of birth", expanded=True):
            pc["NAME"] = st.text_input("Name", value=pc.get("NAME", ""), key=f"pv_name_{active_name}")
            pc["PASSPORT_NO"] = st.text_input("Passport No.", value=pc.get("PASSPORT_NO", ""), key=f"pv_pp_{active_name}")
            pc["DOB"] = st.text_input("DOB (DD/MM/YYYY)", value=pc.get("DOB", ""), key=f"pv_dob_{active_name}")
            pc["HOME_ADDRESS"] = st.text_input("Place of birth", value=pc.get("HOME_ADDRESS", ""), key=f"pv_pob_{active_name}")

        with st.expander("📅 Issue & Expiry dates", expanded=False):
            pc["ISSUE_DATE"] = st.text_input("Issue date", value=pc.get("ISSUE_DATE", ""), key=f"pv_iss_{active_name}")
            pc["EXPIRY_DATE"] = st.text_input("Expiry date", value=pc.get("EXPIRY_DATE", ""), key=f"pv_exp_{active_name}")

        with st.expander("💍 Marital · Children · Religion · Position & more", expanded=False):
            pc["RELATIVE"] = st.text_input("Relative", value=pc.get("RELATIVE", ""), key=f"pv_rel_{active_name}")
            pc["RELIGION"] = st.selectbox("Religion", RELIGION_OPTIONS,
                                           index=0 if pc.get("RELIGION", "MUSLIM") == "MUSLIM" else 1, key=f"pv_relig_{active_name}")
            cur_mar = pc.get("MARITAL_STATUS", "SINGLE")
            pc["MARITAL_STATUS"] = st.selectbox("Marital", ["SINGLE","MARRIED"],
                                                 index=0 if cur_mar == "SINGLE" else 1, key=f"pv_mar_{active_name}")
            ch_opts = ["NIL","1","2","3","4"]; cur_ch = pc.get("NO_OF_CHILDREN", "NIL")
            pc["NO_OF_CHILDREN"] = st.selectbox("Children", ch_opts,
                                                 index=ch_opts.index(cur_ch) if cur_ch in ch_opts else 0, key=f"pv_ch_{active_name}")
            pc["POSITION"] = st.text_input("Position", value=pc.get("POSITION", DEFAULT_POSITION), key=f"pv_pos_{active_name}")
            pc["COUNTRY"] = st.text_input("Country", value=pc.get("COUNTRY", ""), key=f"pv_ctry_{active_name}")
            pc["YEARS_EXP"] = st.text_input("Years exp", value=pc.get("YEARS_EXP", ""), key=f"pv_ye_{active_name}")
            pc["YEARS_EXP2"] = st.text_input("Years exp 2", value=pc.get("YEARS_EXP2", ""), key=f"pv_ye2_{active_name}")

        if st.button("🔄 Update Preview", key=f"pv_update_{active_name}", use_container_width=True, type="primary"):
            new_bytes = _preview_pdf_bytes(pc, cur_agency, cur_level, active_name, files, boxes,
                                            rotation=int(slot.get("rotation", 0)))
            if new_bytes:
                st.session_state[PREVIEW_KEY] = {"file": active_name, "cv": pc, "pdf_bytes": new_bytes}; st.rerun()
            else: st.error("Could not regenerate preview.")

if approve_clicked:
    if pv and pv.get("file") == active_name and pv.get("cv"):
        edited = pv["cv"]
        name_v = edited.get("NAME", "") or name
        passport_v = edited.get("PASSPORT_NO", "") or passport
        dob_v = edited.get("DOB", "") or dob
        issue_v = edited.get("ISSUE_DATE", "") or issue
        expiry_v = edited.get("EXPIRY_DATE", "") or expiry
        pob_v = edited.get("HOME_ADDRESS", "") or pob
        relative_v = edited.get("RELATIVE", "") or relative
        religion_v = edited.get("RELIGION", religion)
        marital_v = edited.get("MARITAL_STATUS", marital)
        children_v = edited.get("NO_OF_CHILDREN", children)
        position_v = edited.get("POSITION", position_value)
        country_v = edited.get("COUNTRY", country_value)
    else:
        name_v, passport_v, dob_v, issue_v, expiry_v = name, passport, dob, issue, expiry
        pob_v, relative_v, religion_v = pob, relative, religion
        marital_v, children_v, position_v, country_v = marital, children, position_value, country_value

    missing_fields = []
    if not name_v.strip(): missing_fields.append("Full name")
    if not passport_v.strip(): missing_fields.append("Passport No.")
    if not dob_v: missing_fields.append("Date of birth")
    if not issue_v: missing_fields.append("Issue date")
    if not expiry_v: missing_fields.append("Expiry date")
    if not pob_v.strip(): missing_fields.append("Place of birth")

    if missing_fields: st.error(f"Please fill in: {', '.join(missing_fields)}")
    else:
        years_str = f"{years_exp} YEARS" if cur_level == "Experienced" else "NIL"
        cv_data = {"NAME": _clean_passport_name(name_v.strip().upper()), "PASSPORT_NO": passport_v.strip().upper(),
                   "ISSUE_DATE": issue_v, "EXPIRY_DATE": expiry_v, "DOB": dob_v, "HOME_ADDRESS": pob_v.strip().upper(),
                   "RELATIVE": relative_v.strip().upper(), "RELIGION": religion_v, "MARITAL_STATUS": marital_v,
                   "NO_OF_CHILDREN": children_v, "AGE": _age_from_dob(dob_v), "POSITION": position_v,
                   "COUNTRY": country_v if cur_level == "Experienced" else "NIL", "YEARS_EXP": years_str,
                   "YEARS_EXP2": years_str if cur_level == "Experienced" else "FIRST TIME"}
        slot["cv_data"] = cv_data; slot["approved"] = True; slot["boxes"] = boxes
        slot["agency"] = cur_agency; slot["level"] = cur_level
        slot["validity_years"] = cur_validity; slot["religion"] = cur_religion
        st.session_state.cv_data = cv_data; st.session_state.approved = True
        st.session_state.pop(PREVIEW_KEY, None)
        next_name = None
        for n in file_names:
            if n != active_name and not per_img[n]["approved"]: next_name = n; break
        if next_name: st.session_state._pending_switch_to = next_name
        st.rerun()

all_approved = all(per_img[n]["approved"] for n in file_names)

if all_approved:
    st.markdown('<hr class="section-divider">', unsafe_allow_html=True)
    st.markdown('<span class="step-tag step-3">STEP 3</span>', unsafe_allow_html=True)
    st.subheader("📄 Generate all CVs")

    _passport_map = {}; _dups = []
    for n in file_names:
        pn = ((per_img[n].get("cv_data") or {}).get("PASSPORT_NO") or "").strip().upper()
        if not pn: continue
        if pn in _passport_map: _dups.append((_passport_map[pn], n, pn))
        else: _passport_map[pn] = n
    if _dups:
        for a, b, pn in _dups:
            st.warning(f"⚠️ **Duplicate passport number `{pn}`** appears in **{a}** and **{b}** — please verify before generating.")

    if st.button("📥 Generate & Download ALL", type="primary", use_container_width=True):
        date_root = desktop_root() / f"output_{datetime.today():%Y-%m-%d}"
        try:
            with st.spinner(f"Generating {len(file_names)} CV(s)…"):
                for n in file_names:
                    s = per_img[n]; cv_data = s["cv_data"]; sboxes = s["boxes"]
                    s_agency = s["agency"]; s_level = s["level"]
                    s_rot = int(s.get("rotation", 0))
                    f_obj = next(f for f in files if f.name == n)
                    simg = Image.open(BytesIO(f_obj.getvalue())).convert("RGB")
                    if s_rot: simg = simg.rotate(-s_rot, expand=True)
                    passport_crop_i = simg.crop(sboxes["passport"])
                    p3x4_crop_i = simg.crop(sboxes["3x4"])
                    full_crop_i = simg.crop(sboxes["full"])
                    cand_name = safe_name(cv_data["NAME"])
                    agencies_i = REAL_AGENCIES if s_agency == BOTH_LABEL else [s_agency]
                    templates_i = {a: find_template(a, s_level) for a in agencies_i}
                    pdf_only_i = set()
                    if s_agency == BOTH_LABEL: pdf_only_i = {a for a in agencies_i if a in PDF_ONLY_AGENCIES_WHEN_BOTH}
                    for a, t in templates_i.items():
                        if t is None: continue
                        folder = date_root / AGENCY_FOLDER[a] / cand_name
                        folder.mkdir(parents=True, exist_ok=True)
                        if a in pdf_only_i:
                            with tempfile.TemporaryDirectory() as td:
                                td_path = Path(td)
                                imgs = {"IMAGE_PASSPORT": td_path/"passport.jpg","IMAGE_3X4": td_path/"3x4.jpg","IMAGE_FULL": td_path/"full.jpg"}
                                passport_crop_i.save(imgs["IMAGE_PASSPORT"], quality=95)
                                p3x4_crop_i.save(imgs["IMAGE_3X4"], quality=95)
                                full_crop_i.save(imgs["IMAGE_FULL"], quality=95)
                                docx_tmp = td_path / f"{cand_name}.docx"
                                pdf_path = folder / f"{cand_name}.pdf"
                                fill_cv(t, cv_data, imgs, docx_tmp, pdf_path)
                        else:
                            passport_path = folder/"passport.jpg"; photo3_path = folder/"3x4.jpg"; full_path = folder/"full.jpg"
                            passport_crop_i.save(passport_path, quality=95, optimize=True)
                            p3x4_crop_i.save(photo3_path, quality=95, optimize=True)
                            full_crop_i.save(full_path, quality=95, optimize=True)
                            docx_path = folder / f"{cand_name}.docx"; pdf_path = folder / f"{cand_name}.pdf"
                            imgs = {"IMAGE_PASSPORT": passport_path,"IMAGE_3X4": photo3_path,"IMAGE_FULL": full_path}
                            fill_cv(t, cv_data, imgs, docx_path, pdf_path)
                notes_lines = []
                for n in file_names:
                    note = (per_img[n].get("notes") or "").strip()
                    if note: notes_lines.append(f"{n}: {note}")
                if notes_lines: (date_root / "notes.txt").write_text("\n".join(notes_lines), encoding="utf-8")
            zip_buffer = io.BytesIO()
            with zipfile.ZipFile(zip_buffer, "w", compression=zipfile.ZIP_DEFLATED) as zf:
                for item in date_root.rglob("*"):
                    if item.is_file(): zf.write(item, item.relative_to(date_root))
            zip_buffer.seek(0)
            st.success(f"✅ Generated {len(file_names)} CV(s). Saved to: {date_root}")
            st.download_button("📦 Download EVERYTHING (ZIP)", data=zip_buffer.getvalue(),
                                file_name=f"All_CVs_{datetime.today():%Y-%m-%d}.zip",
                                mime="application/zip", type="primary", use_container_width=True)
        except Exception as e:
            st.error(f"Generation failed: {e}")