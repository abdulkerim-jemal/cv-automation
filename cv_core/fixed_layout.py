from pathlib import Path
import os, subprocess, tempfile, shutil
from io import BytesIO
from PIL import Image
from docx import Document
from docx.shared import Inches

FIXED={"Asail":{"passport":(0.0,4.80,3.80,5.00),"3x4":(0.55,0.50,1.60,1.38),"full":(0.35,5.00,2.35,4.65)},"Al Zaid":{"passport":(0.55,1.45,3.35,4.65),"3x4":(0.55,0.50,1.60,1.38),"full":(0.35,5.00,2.35,4.65)}}


def _iter_paragraphs(parent):
 for paragraph in parent.paragraphs:
  yield paragraph
 for table in parent.tables:
  for row in table.rows:
   for cell in row.cells:
    yield from _iter_paragraphs(cell)


def _replace_token(paragraph, token, value):
 """Replace a token without removing the template's run-level formatting."""
 while token in "".join(run.text for run in paragraph.runs):
  all_text="".join(run.text for run in paragraph.runs)
  start=all_text.index(token); end=start+len(token)
  cursor=0; first=last=None
  for index,run in enumerate(paragraph.runs):
   next_cursor=cursor+len(run.text)
   if first is None and start < next_cursor: first=(index,start-cursor)
   if end <= next_cursor:
    last=(index,end-cursor); break
   cursor=next_cursor
  if first is None or last is None: return
  first_index,first_offset=first; last_index,last_offset=last
  if first_index==last_index:
   run=paragraph.runs[first_index]
   run.text=run.text[:first_offset]+str(value)+run.text[last_offset:]
  else:
   first_run=paragraph.runs[first_index]
   last_run=paragraph.runs[last_index]
   first_run.text=first_run.text[:first_offset]+str(value)
   for index in range(first_index+1,last_index): paragraph.runs[index].text=""
   last_run.text=last_run.text[last_offset:]


def replace_text_only(template,data,out_docx):
 d=Document(str(template))
 parents=[d]
 for section in d.sections:
  parents.extend([section.header,section.footer,section.first_page_header,section.first_page_footer])
 for parent in parents:
  for paragraph in _iter_paragraphs(parent):
   if "{{" not in paragraph.text: continue
   for key,value in data.items(): _replace_token(paragraph,"{{%s}}"%key,value)
 d.save(str(out_docx))

def _find_soffice():
 found=(shutil.which("soffice") or shutil.which("soffice.exe") or
        shutil.which("soffice.com") or shutil.which("libreoffice"))
 if found: return found
 candidates=[
  r"C:\Program Files\LibreOffice\program\soffice.exe",
  r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
  str(Path(os.environ.get("LOCALAPPDATA", "")) / "Programs" / "LibreOffice" / "program" / "soffice.exe"),
  "/Applications/LibreOffice.app/Contents/MacOS/soffice",
  "/usr/bin/soffice","/usr/local/bin/soffice","/snap/bin/libreoffice",
 ]
 for c in candidates:
  if Path(c).exists(): return c
 # Windows sometimes records a custom LibreOffice install folder in the registry.
 try:
  import winreg
  for hive in (winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER):
   for key_path in (r"SOFTWARE\LibreOffice\LibreOffice", r"SOFTWARE\WOW6432Node\LibreOffice\LibreOffice"):
    try:
     with winreg.OpenKey(hive, key_path) as key:
      install_dir,_=winreg.QueryValueEx(key,"Path")
      candidate=Path(install_dir) / "program" / "soffice.exe"
      if candidate.exists(): return str(candidate)
    except OSError: pass
 except ImportError: pass
 return None

def _convert_docx(docx,pdf):
    # Preferred: LibreOffice is free and works without Microsoft Word.
    soffice=_find_soffice()
    if soffice:
        with tempfile.TemporaryDirectory() as td:
            for export_format in ("pdf:writer_pdf_Export", "pdf"):
                subprocess.run(
                    [soffice,"--headless","--norestore","--convert-to",export_format,
                     "--outdir",td,str(docx)],
                    capture_output=True,timeout=120
                )
                src=Path(td)/(Path(docx).stem+".pdf")
                if src.exists() and src.stat().st_size > 0:
                    shutil.copyfile(src,pdf)
                    return

    # Fallback on Windows/macOS when Microsoft Word is installed.
    try:
        from docx2pdf import convert
        convert(str(docx),str(pdf))
        if Path(pdf).exists() and Path(pdf).stat().st_size > 0:
            return
    except Exception:
        pass

    raise RuntimeError(
        "PDF export needs LibreOffice or Microsoft Word. "
        "Install LibreOffice (recommended) or Microsoft Word, then run the app again."
    )


def _render_portable_fallback_pdf(data, images, agency, out_pdf):
 """Create a styled agency CV when no local DOCX-to-PDF engine is available."""
 import fitz
 from .image_processing import prepare_cv_image

 doc=fitz.open()
 page1=doc.new_page(width=595, height=842)
 page2=doc.new_page(width=595, height=842)
 blue=(0.02,0.65,0.85); red=(0.93,0.02,0.02); black=(0.05,0.05,0.05); border=(0.15,0.15,0.15)

 def add_image(page, path, rect, anchor="center"):
  width=max(1, int(rect.width*2))
  height=max(1, int(rect.height*2))
  prepared=prepare_cv_image(Image.open(path), width, height, mode="contain" if anchor=="contain" else "cover", anchor="center" if anchor=="contain" else anchor)
  stream=BytesIO(); prepared.save(stream, "JPEG", quality=95)
  page.insert_image(rect, stream=stream.getvalue(), keep_proportion=False, overlay=True)

 def section(page, rect, title):
  page.draw_rect(rect, color=border, width=0.7)
  page.draw_rect(fitz.Rect(rect.x0,rect.y0,rect.x1,rect.y0+22), color=blue, fill=blue)
  page.insert_text(fitz.Point(rect.x0+8,rect.y0+16),title,fontsize=11,fontname="hebo",color=red)

 def table(page, rect, rows, religion_red=False):
  row_h=(rect.height-22)/max(1,len(rows))
  split=rect.x0+rect.width*.42
  for index,(label,value) in enumerate(rows):
   y=rect.y0+22+index*row_h
   page.draw_rect(fitz.Rect(rect.x0,y,rect.x1,y+row_h),color=border,width=0.45)
   page.draw_line(fitz.Point(split,y),fitz.Point(split,y+row_h),color=border,width=0.45)
   page.insert_text(fitz.Point(rect.x0+7,y+row_h*.68),label,fontsize=8.8,fontname="hebo",color=black)
   value_color=red if religion_red and label=="RELIGION" else black
   value_font="hebo" if value_color==red else "helv"
   page.insert_text(fitz.Point(split+7,y+row_h*.68),str(value or "NIL"),fontsize=8.8,fontname=value_font,color=value_color)

 page1.draw_rect(fitz.Rect(6,6,589,836),color=(0.38,0.10,0.60),width=1.8)
 page1.insert_textbox(fitz.Rect(45,25,550,65),f"{agency.upper()} FOREIGN EMPLOYMENT AGENT\nETHIOPIAN HOUSE MAID",fontsize=16,fontname="hebo",align=1,color=black)
 page1.draw_rect(fitz.Rect(42,85,173,225),color=border,width=0.7)
 add_image(page1,images["IMAGE_3X4"],fitz.Rect(47,90,168,220),"contain")
 section(page1,fitz.Rect(205,85,555,220),"CANDIDATE DETAILS")
 table(page1,fitz.Rect(205,85,555,220),[
  ("NAME IN FULL",data.get("NAME","")),("HOME ADDRESS",data.get("HOME_ADDRESS","")),
  ("RELATIVE",data.get("RELATIVE","")),("POSITION APPLIED",data.get("POSITION","")),
  ("CONTRACT PERIOD",data.get("YEARS_EXP2","")),
 ])
 section(page1,fitz.Rect(20,250,265,350),"PASSPORT DETAILS")
 table(page1,fitz.Rect(20,250,265,350),[
  ("PASSPORT NO.",data.get("PASSPORT_NO","")),("DATE OF ISSUE",data.get("ISSUE_DATE","")),
  ("DATE OF EXPIRY",data.get("EXPIRY_DATE","")),
 ])
 section(page1,fitz.Rect(290,250,555,430),"PERSONAL INFORMATION")
 table(page1,fitz.Rect(290,250,555,430),[
  ("NATIONALITY","ETHIOPIA"),("RELIGION",data.get("RELIGION","")),
  ("DATE OF BIRTH",data.get("DOB","")),("AGE",data.get("AGE","")),
  ("PLACE OF BIRTH",data.get("HOME_ADDRESS","")),("MARITAL STATUS",data.get("MARITAL_STATUS","")),
  ("NO. OF CHILDREN",data.get("NO_OF_CHILDREN","")),
 ],religion_red=True)
 page1.draw_rect(fitz.Rect(34,380,265,790),color=border,width=0.7)
 add_image(page1,images["IMAGE_FULL"],fitz.Rect(40,386,259,784),"contain")
 section(page1,fitz.Rect(290,460,555,565),"EDUCATIONAL BACKGROUND")
 table(page1,fitz.Rect(290,460,555,565),[("CERTIFICATE OF COMPETENCY","DOMESTIC WORK"),("ELEMENTRY","YES"),("HIGH SCHOOL","NO"),("COLLEGE","NO")])
 section(page1,fitz.Rect(290,590,555,690),"LANGUAGES SPEAKING")
 table(page1,fitz.Rect(290,590,555,690),[("ENGLISH","POOR"),("ARABIC","POOR"),("AMHARIC","GOOD"),("OROMIYA","GOOD")])
 section(page1,fitz.Rect(290,715,555,790),"WORK EXPERIENCE")
 table(page1,fitz.Rect(290,715,555,790),[("POSITION",data.get("POSITION","")),("COUNTRY",data.get("COUNTRY","")),("EXPERIENCE",data.get("YEARS_EXP",""))])

 page2.draw_rect(fitz.Rect(6,6,589,836),color=(0.38,0.10,0.60),width=1.8)
 page2.insert_text(fitz.Point(30,35),"PASSPORT COPY",fontsize=14,fontname="hebo",color=black)
 page2.draw_rect(fitz.Rect(15,50,350,550),color=border,width=0.7)
 add_image(page2,images["IMAGE_PASSPORT"],fitz.Rect(20,55,345,545),"contain")
 doc.save(str(out_pdf),garbage=4,deflate=True)
 doc.close()

def render_fixed_pdf(template,data,images,agency,out_pdf):
 import fitz
 from .image_processing import prepare_cv_image
 with tempfile.TemporaryDirectory() as td:
  bd=Path(td)/"base.docx"; bp=Path(td)/"base.pdf"; replace_text_only(template,data,bd)
  try: _convert_docx(bd,bp)
  except RuntimeError:
   _render_portable_fallback_pdf(data,images,agency,out_pdf)
   return
  src=fitz.open(bp)
  if src.page_count<2: raise RuntimeError(f"Template rendered only {src.page_count} pages")
  out=fitz.open(); out.insert_pdf(src,from_page=0,to_page=1); cfg=FIXED[agency]
  def put(page,key,path):
   x,y,w,h=cfg[key]; rect=fitz.Rect(x*72,y*72,(x+w)*72,(y+h)*72)
   page.add_redact_annot(rect,fill=(1,1,1))
   page.apply_redactions()
   mode="contain"
   im=prepare_cv_image(Image.open(path),int(w*150),int(h*150),mode=mode,anchor="center")
   tmp=Path(td)/(key+".jpg"); im.save(tmp,"JPEG",quality=95); page.insert_image(rect,filename=str(tmp),keep_proportion=False,overlay=True)
  put(out[0],"full",images["IMAGE_FULL"]); put(out[1],"3x4",images["IMAGE_3X4"]); put(out[1],"passport",images["IMAGE_PASSPORT"])
  out.save(str(out_pdf),garbage=4,deflate=True); out.close(); src.close()

def make_visual_docx_from_pdf(pdf_path,out_docx):
 import fitz
 with tempfile.TemporaryDirectory() as td:
  pdf=fitz.open(str(pdf_path)); d=Document(); d.add_paragraph(); s=d.sections[0]; s.page_width=Inches(8.5); s.page_height=Inches(11); s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(0)
  for i,page in enumerate(pdf):
   if i: d.add_page_break()
   pix=page.get_pixmap(matrix=fitz.Matrix(2,2),alpha=False); img=Path(td)/f"p{i}.png"; pix.save(str(img)); p=d.paragraphs[-1]; p.paragraph_format.space_before=0; p.paragraph_format.space_after=0; p.add_run().add_picture(str(img),width=Inches(8.5),height=Inches(11))
  d.save(str(out_docx)); pdf.close()
