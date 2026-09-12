from __future__ import annotations

import copy
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, RGBColor

from .image_processing import prepare_cv_image

PHOTO_BOX_IN={"IMAGE_3X4":(1.52,1.30),"IMAGE_FULL":(2.27,4.03),"IMAGE_PASSPORT":(3.35,4.65)}
BOLD_RED_KEYS={"RELIGION","YEARS_EXP2"}


def iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from iter_table_paragraphs(nested)


def iter_all_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        yield from iter_table_paragraphs(table)


def _clear_runs(p):
    for run in list(p.runs):
        run._element.getparent().remove(run._element)


def _cell_dimensions(p):
    tc=p._p.getparent()
    while tc is not None and tc.tag!=qn("w:tc"): tc=tc.getparent()
    if tc is None:return None
    tcPr=tc.find(qn("w:tcPr")); width=None
    if tcPr is not None:
        tcW=tcPr.find(qn("w:tcW"))
        if tcW is not None and tcW.get(qn("w:w")):
            try: width=int(tcW.get(qn("w:w")))/1440
            except: pass
    row=tc.getparent(); height=None
    if row is not None:
        trPr=row.find(qn("w:trPr"))
        if trPr is not None:
            trH=trPr.find(qn("w:trHeight"))
            if trH is not None and trH.get(qn("w:val")):
                try: height=int(trH.get(qn("w:val")))/1440
                except: pass
    return (width,height) if width and height else None


def _lock_cell_exact(p,w_in,h_in):
    tc=p._p.getparent()
    while tc is not None and tc.tag!=qn("w:tc"): tc=tc.getparent()
    if tc is None:return
    tcPr=tc.find(qn("w:tcPr"))
    if tcPr is None: tcPr=OxmlElement("w:tcPr"); tc.insert(0,tcPr)
    tcW=tcPr.find(qn("w:tcW"))
    if tcW is None: tcW=OxmlElement("w:tcW"); tcPr.append(tcW)
    tcW.set(qn("w:type"),"dxa"); tcW.set(qn("w:w"),str(int(round(w_in*1440))))
    tr=tc.getparent()
    if tr is not None and tr.tag==qn("w:tr"):
        trPr=tr.find(qn("w:trPr"))
        if trPr is None: trPr=OxmlElement("w:trPr"); tr.insert(0,trPr)
        trH=trPr.find(qn("w:trHeight"))
        if trH is None: trH=OxmlElement("w:trHeight"); trPr.append(trH)
        trH.set(qn("w:hRule"),"exact"); trH.set(qn("w:val"),str(int(round(h_in*1440))))
    tbl=tc.getparent()
    while tbl is not None and tbl.tag!=qn("w:tbl"): tbl=tbl.getparent()
    if tbl is not None:
        tblPr=tbl.find(qn("w:tblPr"))
        if tblPr is None: tblPr=OxmlElement("w:tblPr"); tbl.insert(0,tblPr)
        layout=tblPr.find(qn("w:tblLayout"))
        if layout is None: layout=OxmlElement("w:tblLayout"); tblPr.append(layout)
        layout.set(qn("w:type"),"fixed")


def _paragraph_box(p,key):
    dims=_cell_dimensions(p)
    if dims:
        w,h=dims
        return max(0.4,w-0.04),max(0.4,h-0.04)
    return PHOTO_BOX_IN.get(key,(2,2.5))


def _set_paragraph_tight(p):
    pf=p.paragraph_format
    pf.space_before=0; pf.space_after=0
    pf.line_spacing=__import__('docx').shared.Pt(1)
    pPr=p._p.get_or_add_pPr()
    spacing=pPr.find(qn('w:spacing'))
    if spacing is None: spacing=OxmlElement('w:spacing'); pPr.append(spacing)
    spacing.set(qn('w:before'),'0'); spacing.set(qn('w:after'),'0'); spacing.set(qn('w:line'),'20'); spacing.set(qn('w:lineRule'),'exact')


def _inline_to_anchor(inline, left_in, top_in, behind=False):
    # LibreOffice is most reliable when anchors match the native Word anchors
    # already present in these templates. Build one from the inline picture and
    # use the standard Word anchor schema.
    anchor=OxmlElement("wp:anchor")
    for k,v in {"distT":"0","distB":"0","distL":"0","distR":"0",
                "simplePos":"0","relativeHeight":"251659264",
                "behindDoc":"1" if behind else "0","locked":"0",
                "layoutInCell":"1","allowOverlap":"1"}.items():
        anchor.set(qn("wp:"+k),v)
    simple=OxmlElement("wp:simplePos"); simple.set("x","0"); simple.set("y","0"); anchor.append(simple)
    ph=OxmlElement("wp:positionH"); ph.set("relativeFrom","page")
    off=OxmlElement("wp:posOffset"); off.text=str(int(left_in*914400)); ph.append(off); anchor.append(ph)
    pv=OxmlElement("wp:positionV"); pv.set("relativeFrom","page")
    off=OxmlElement("wp:posOffset"); off.text=str(int(top_in*914400)); pv.append(off); anchor.append(pv)
    anchor.append(OxmlElement("wp:wrapNone"))
    extent=inline.find(qn("wp:extent"))
    if extent is not None: anchor.append(copy.deepcopy(extent))
    effect=inline.find(qn("wp:effectExtent"))
    if effect is not None: anchor.append(copy.deepcopy(effect))
    docpr=inline.find(qn("wp:docPr"))
    if docpr is not None:
        docpr=copy.deepcopy(docpr); docpr.set("id","9001"); docpr.set("name","Passport Photo"); anchor.append(docpr)
    cnv=inline.find(qn("wp:cNvGraphicFramePr"))
    if cnv is not None: anchor.append(copy.deepcopy(cnv))
    graphic=inline.find(qn("a:graphic"))
    if graphic is not None: anchor.append(copy.deepcopy(graphic))
    return anchor

def add_floating_picture(paragraph,path,width_in,height_in,left_in=0.55,top_in=2.05):
    _clear_runs(paragraph); _set_paragraph_tight(paragraph)
    run=paragraph.add_run()
    inline=run.add_picture(str(path),width=Inches(width_in),height=Inches(height_in))._inline
    anchor=_inline_to_anchor(inline,left_in,top_in)
    inline.getparent().replace(inline,anchor)
    return anchor


def _process_image_placeholder(p,key,path):
    _clear_runs(p); _set_paragraph_tight(p)
    run=p.add_run()
    w,h=_paragraph_box(p,key)
    dpi=120
    target=(max(1,int(w*dpi)),max(1,int(h*dpi)))
    prepared=prepare_cv_image(Image.open(path),*target,mode="contain",anchor="center")
    tmp=Path(tempfile.mkstemp(suffix=".jpg")[1]); prepared.save(tmp,"JPEG",quality=95)
    try:
        run.add_picture(str(tmp),width=Inches(w),height=Inches(h))
        if key == "IMAGE_FULL":
            spacer=p.add_run("\u200b")
            spacer.font.size=__import__('docx').shared.Pt(10)
            spacer.font.color.rgb=RGBColor(255,255,255)
    finally:
        try: tmp.unlink()
        except: pass
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    dims=_cell_dimensions(p)
    if dims:
        _lock_cell_exact(p,dims[0],dims[1])
    else:
        _lock_cell_exact(p,w,h)


def render_cv_data(doc,data):
    for p in list(iter_all_paragraphs(doc)):
        text=p.text
        if not text or "{{" not in text: continue
        # Passport is handled separately below.
        if "{{IMAGE_PASSPORT}}" in text: continue
        for key,val in data.items():
            ph="{{%s}}"%key
            if ph in text:
                first=p.runs[0] if p.runs else None
                new=text.replace(ph,str(val))
                _clear_runs(p); r=p.add_run(new)
                if first is not None and first._element.rPr is not None: r._element.insert(0,copy.deepcopy(first._element.rPr))
                if key in BOLD_RED_KEYS:
                    r.bold=True; r.font.color.rgb=RGBColor(255,0,0)
                break


def place_passport_in_second_page_header(doc, path):
    section=doc.sections[0]
    section.different_first_page_header_footer=True
    # Keep first-page header empty; default header is used from page 2 onward.
    first=section.first_page_header
    if not first.paragraphs:
        first.add_paragraph()
    default=section.header
    section.header_distance=__import__('docx').shared.Inches(0.02)
    p=default.paragraphs[0] if default.paragraphs else default.add_paragraph()
    add_floating_picture(p,path,*PHOTO_BOX_IN["IMAGE_PASSPORT"],left_in=0.63,top_in=1.50)


def place_cv_images(doc,images):
    for p in list(iter_all_paragraphs(doc)):
        text=p.text
        if "{{IMAGE_PASSPORT}}" in text:
            _clear_runs(p)
            continue
        for key,path in images.items():
            if key in ("IMAGE_PASSPORT",): continue
            if "{{%s}}"%key in text:
                _process_image_placeholder(p,key,path)
                break


def force_calibri(doc):
    for p in iter_all_paragraphs(doc):
        for r in p.runs:
            r.font.name="Calibri"
            rPr=r._element.get_or_add_rPr(); rFonts=rPr.find(qn("w:rFonts"))
            if rFonts is None: rFonts=OxmlElement("w:rFonts"); rPr.append(rFonts)
            for a in ("w:ascii","w:hAnsi","w:eastAsia","w:cs"): rFonts.set(qn(a),"Calibri")


def find_soffice():
    for n in ("soffice","libreoffice","soffice.exe"):
        p=shutil.which(n)
        if p:return p
    for p in (r"C:\Program Files\LibreOffice\program\soffice.exe",r"C:\Program Files (x86)\LibreOffice\program\soffice.exe","/Applications/LibreOffice.app/Contents/MacOS/soffice","/usr/bin/soffice","/usr/lib/libreoffice/program/soffice"):
        if os.path.exists(p):return p
    return None


def convert_docx_to_pdf(docx_path,pdf_path):
    soffice=find_soffice()
    if soffice:
        with tempfile.TemporaryDirectory() as td:
            r=subprocess.run([soffice,"--headless","--norestore","--convert-to","pdf","--outdir",td,str(docx_path)],capture_output=True,text=True,timeout=120)
            produced=Path(td)/f"{Path(docx_path).stem}.pdf"
            if produced.exists() and produced.stat().st_size>0:
                Path(pdf_path).parent.mkdir(parents=True,exist_ok=True); shutil.copyfile(produced,pdf_path); return True
    try:
        from docx2pdf import convert
        convert(str(docx_path),str(pdf_path)); return Path(pdf_path).exists() and Path(pdf_path).stat().st_size>0
    except Exception:return False


def validate_docx_placeholders(doc):
    leftover=[]
    for p in iter_all_paragraphs(doc):
        if "{{" in p.text:leftover.append(p.text)
    return leftover


def pdf_page_count(path):
    try:
        import fitz
        d=fitz.open(str(path)); n=d.page_count; d.close(); return n
    except Exception:return None


def validate_output(docx_path,pdf_path,expected_pages=2):
    doc=Document(str(docx_path)); leftovers=validate_docx_placeholders(doc)
    pages=pdf_page_count(pdf_path) if Path(pdf_path).exists() else None
    return {"placeholders":leftovers,"pdf_pages":pages,"expected_pages":expected_pages,"ok":not leftovers and pages==expected_pages}


def fill_cv(template_path,data,images,docx_out,pdf_out,expected_pages=2):
    doc=Document(str(template_path))
    render_cv_data(doc,data)
    place_passport_in_second_page_header(doc, images["IMAGE_PASSPORT"])
    place_cv_images(doc,images)
    force_calibri(doc)
    doc.save(str(docx_out))
    made=convert_docx_to_pdf(docx_out,pdf_out)
    validation=validate_output(docx_out,pdf_out,expected_pages) if made else {"placeholders":validate_docx_placeholders(doc),"pdf_pages":None,"expected_pages":expected_pages,"ok":False}
    if not validation["ok"]:
        raise RuntimeError(f"CV layout validation failed: {validation}")
    return validation
